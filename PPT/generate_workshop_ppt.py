#!/usr/bin/env python3
"""生成“从零开始搭建 Agent”Workshop 演示文稿。

本脚本只依赖 python-pptx 与 Pillow。所有主视觉均使用 PowerPoint
原生对象，PNG 仅用于生成后的视觉检查，不会被铺成 PPT 页面背景。
"""

from __future__ import annotations

import json
import math
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = Path(__file__).resolve().parent
ASSET_DIR = OUT_DIR / "assets"
PREVIEW_DIR = OUT_DIR / "previews"
RENDER_DIR = OUT_DIR / "rendered"
PPTX_PATH = OUT_DIR / "agent_from_scratch_workshop.pptx"
NOTES_PATH = OUT_DIR / "speaker_notes.md"
REPORT_PATH = OUT_DIR / "qa_report.json"

SLIDE_W = 13.333
SLIDE_H = 7.5


class C:
    """全局配色；页面不得在此之外随机创建颜色。"""

    NAVY = "17324D"
    BLUE = "315B7D"
    ORANGE = "D97706"
    GREEN = "2F855A"
    RED = "C2413B"
    INK = "263746"
    GRAY = "64748B"
    MID = "94A3B8"
    LIGHT = "E2E8F0"
    PALE = "F4F7FA"
    WHITE = "FFFFFF"
    PALE_BLUE = "EAF1F7"
    PALE_ORANGE = "FFF4E6"
    PALE_GREEN = "EAF6EF"
    PALE_RED = "FCEDEC"


FONT_CN = "Noto Sans CJK SC"
FONT_EN = "Aptos"
FONT_MONO = "DejaVu Sans Mono"


@dataclass
class SlideMeta:
    """保存页面 ID、标题、阶段与讲者备注。"""

    slide_id: str
    title: str
    stage: str
    minutes: str
    message: str
    notes: list[str] = field(default_factory=list)
    transition: str = ""


@dataclass
class ShapeRecord:
    """记录元素边界，用于生成后检查。"""

    slide_no: int
    name: str
    left: float
    top: float
    width: float
    height: float
    allow_overlap: bool = False

    @property
    def right(self) -> float:
        return self.left + self.width

    @property
    def bottom(self) -> float:
        return self.top + self.height


def rgb(hex_color: str) -> RGBColor:
    """把六位十六进制颜色转换为 python-pptx 颜色。"""

    return RGBColor.from_string(hex_color)


def set_run_font(run, name: str, size: float, color: str, bold: bool = False) -> None:
    """同时设置拉丁字体与东亚字体，避免中文在 PowerPoint 中随机回退。"""

    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    east_asia = rpr.find(qn("a:ea"))
    if east_asia is None:
        east_asia = OxmlElement("a:ea")
        rpr.append(east_asia)
    east_asia.set("typeface", name)


def choose_font(text: str) -> str:
    """纯英文使用 Aptos；含中文或其他东亚字符时使用统一中文字体。"""

    return FONT_CN if any(ord(char) > 127 for char in text) else FONT_EN


def hexagon_points(x: float, y: float, w: float, h: float) -> list[tuple[float, float]]:
    """返回六边形顶点，供预览图使用。"""

    return [(x + w * .2, y), (x + w * .8, y), (x + w, y + h / 2),
            (x + w * .8, y + h), (x + w * .2, y + h), (x, y + h / 2)]


class DeckBuilder:
    """统一管理 PPT 原生对象、预览对象、备注和质量检查。"""

    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]
        self.slide = None
        self.meta: list[SlideMeta] = []
        self.records: list[ShapeRecord] = []
        self.preview_ops: list[list[dict]] = []
        self.current_no = 0

    def new_slide(self, meta: SlideMeta, show_progress: bool = True) -> None:
        """创建空白页，并添加统一标题、Stage 进度条和页脚。"""

        self.slide = self.prs.slides.add_slide(self.blank)
        self.current_no = len(self.prs.slides)
        self.meta.append(meta)
        self.preview_ops.append([])
        self.rect(0, 0, SLIDE_W, SLIDE_H, C.WHITE, C.WHITE, name="background", allow_overlap=True)
        if meta.slide_id != "cover":
            self.text(meta.title, .68, .42, 11.9, .48, 31, C.NAVY, bold=True,
                      name="title")
            self.line(.68, 1.08, 12.65, 1.08, C.LIGHT, 1.2, name="title_rule")
            if show_progress and meta.stage.startswith("S"):
                self.progress(meta.stage)
            self.footer(meta.stage)

    def _record(self, name: str, x: float, y: float, w: float, h: float,
                allow_overlap: bool = False) -> None:
        self.records.append(ShapeRecord(self.current_no, name, x, y, w, h, allow_overlap))

    def _op(self, kind: str, **kwargs) -> None:
        kwargs["kind"] = kind
        self.preview_ops[-1].append(kwargs)

    def text(self, text: str, x: float, y: float, w: float, h: float, size: float = 20,
             color: str = C.INK, bold: bool = False, align=PP_ALIGN.LEFT,
             valign=MSO_ANCHOR.MIDDLE, font: str | None = None, margin: float = .05,
             name: str = "text", allow_overlap: bool = False) -> None:
        """添加无边框文本框；换行由内容显式控制。"""

        shape = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = f"{self.current_no:02d}_{name}"
        tf = shape.text_frame
        tf.clear()
        tf.margin_left = tf.margin_right = Inches(margin)
        tf.margin_top = tf.margin_bottom = Inches(.01)
        tf.vertical_anchor = valign
        font_name = font or choose_font(text)
        for idx, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = line
            p.alignment = align
            p.space_after = Pt(0)
            p.space_before = Pt(0)
            p.line_spacing = 1.05
            for run in p.runs:
                set_run_font(run, font_name, size, color, bold)
        self._record(name, x, y, w, h, allow_overlap)
        self._op("text", text=text, x=x, y=y, w=w, h=h, size=size, color=color,
                 bold=bold, align=str(align))

    def rect(self, x: float, y: float, w: float, h: float, fill: str = C.WHITE,
             line: str = C.LIGHT, radius: bool = False, line_width: float = 1,
             name: str = "rect", allow_overlap: bool = False, transparency: int = 0):
        """添加矩形或圆角矩形。"""

        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = self.slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = f"{self.current_no:02d}_{name}"
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.fill.transparency = transparency
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
        self._record(name, x, y, w, h, allow_overlap)
        self._op("roundrect" if radius else "rect", x=x, y=y, w=w, h=h,
                 fill=fill, line=line, line_width=line_width)
        return shape

    def box(self, text: str, x: float, y: float, w: float, h: float,
            fill: str = C.WHITE, line: str = C.LIGHT, size: float = 18,
            color: str = C.INK, bold: bool = False, radius: bool = True,
            name: str = "box", allow_overlap: bool = False) -> None:
        """添加带居中文字的卡片。"""

        self.rect(x, y, w, h, fill, line, radius, name=f"{name}_bg",
                  allow_overlap=allow_overlap)
        self.text(text, x + .08, y + .06, w - .16, h - .12, size, color, bold,
                  PP_ALIGN.CENTER, name=f"{name}_text", allow_overlap=True)

    def hex(self, text: str, x: float, y: float, w: float, h: float,
            fill: str = C.PALE_ORANGE, line: str = C.ORANGE, size: float = 17,
            name: str = "tool", text_color: str = C.INK) -> None:
        """添加统一的工具六边形。"""

        shape = self.slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = f"{self.current_no:02d}_{name}"
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
        shape.line.color.rgb = rgb(line); shape.line.width = Pt(1.2)
        self._record(name, x, y, w, h)
        self._op("hex", text=text, x=x, y=y, w=w, h=h, fill=fill, line=line,
                 size=size, color=text_color)
        tf = shape.text_frame; tf.clear(); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.text = text; p.alignment = PP_ALIGN.CENTER
        for run in p.runs: set_run_font(run, choose_font(text), size, text_color, True)

    def line(self, x1: float, y1: float, x2: float, y2: float, color: str = C.MID,
             width: float = 1.5, dash: bool = False, name: str = "line") -> None:
        """添加直线；线条允许与其他元素相交。"""

        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        shape.name = f"{self.current_no:02d}_{name}"
        shape.line.color.rgb = rgb(color); shape.line.width = Pt(width)
        if dash: shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        self._record(name, min(x1, x2), min(y1, y2), abs(x2-x1) or .01,
                     abs(y2-y1) or .01, True)
        self._op("line", x1=x1, y1=y1, x2=x2, y2=y2, color=color, width=width,
                 dash=dash)

    def arrow(self, x: float, y: float, w: float, h: float, color: str = C.BLUE,
              direction: str = "right", name: str = "arrow") -> None:
        """添加块状箭头，保证在 PowerPoint 中可编辑。"""

        shape_type = {
            "right": MSO_SHAPE.RIGHT_ARROW,
            "left": MSO_SHAPE.LEFT_ARROW,
            "down": MSO_SHAPE.DOWN_ARROW,
            "up": MSO_SHAPE.UP_ARROW,
        }[direction]
        shape = self.slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = f"{self.current_no:02d}_{name}"
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(color)
        shape.line.color.rgb = rgb(color)
        self._record(name, x, y, w, h, True)
        self._op("arrow", x=x, y=y, w=w, h=h, color=color, direction=direction)

    def chip(self, text: str, x: float, y: float, w: float, fill: str = C.PALE_BLUE,
             color: str = C.NAVY, name: str = "chip") -> None:
        """添加小型标签。"""

        self.rect(x, y, w, .36, fill, fill, True, name=f"{name}_bg", allow_overlap=True)
        self.text(text, x + .03, y + .01, w - .06, .33, 11, color, True,
                  PP_ALIGN.CENTER, name=f"{name}_text", allow_overlap=True)

    def progress(self, stage: str) -> None:
        """添加 S0–S7 的统一阶段进度条。"""

        current = int(stage[1]) if len(stage) > 1 and stage[1].isdigit() else -1
        start_x, y, gap, width = 8.95, .93, .10, .37
        for i in range(8):
            fill = C.NAVY if i == current else (C.MID if i < current else C.LIGHT)
            self.rect(start_x + i * (width + gap), y, width, .08, fill, fill,
                      name=f"progress_{i}", allow_overlap=True)

    def footer(self, stage: str) -> None:
        """添加统一页脚。"""

        self.text("Agent from Scratch", .68, 7.08, 2.4, .18, 9.5, C.GRAY,
                  name="footer_left")
        self.text(stage, 11.65, 7.08, .55, .18, 9.5, C.GRAY, True,
                  PP_ALIGN.RIGHT, name="footer_stage")
        self.text(str(self.current_no), 12.27, 7.08, .38, .18, 9.5, C.GRAY,
                  False, PP_ALIGN.RIGHT, name="footer_no")

    def table(self, data: Sequence[Sequence[str]], x: float, y: float, w: float,
              h: float, col_widths: Sequence[float] | None = None,
              header: bool = True, font_size: float = 15, name: str = "table") -> None:
        """添加统一样式的原生 PowerPoint 表格。"""

        rows, cols = len(data), len(data[0])
        shape = self.slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = f"{self.current_no:02d}_{name}"
        table = shape.table
        if col_widths:
            total = sum(col_widths)
            for idx, value in enumerate(col_widths):
                table.columns[idx].width = Inches(w * value / total)
        for r in range(rows):
            for c in range(cols):
                cell = table.cell(r, c)
                cell.text = str(data[r][c])
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb(C.PALE_BLUE if header and r == 0 else C.WHITE)
                cell.margin_left = cell.margin_right = Inches(.08)
                cell.margin_top = cell.margin_bottom = Inches(.04)
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.LEFT if c else PP_ALIGN.CENTER
                    for run in p.runs:
                        set_run_font(run, choose_font(str(data[r][c])), font_size,
                                     C.NAVY if r == 0 else C.INK, r == 0)
        self._record(name, x, y, w, h)
        self._op("table", data=[list(row) for row in data], x=x, y=y, w=w, h=h,
                 header=header, font_size=font_size)

    def placeholder(self, label: str, path: Path, x: float, y: float, w: float,
                    h: float, name: str) -> None:
        """素材存在时插图；否则生成带文件名的原生占位框。"""

        if path.exists():
            self.slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
            self._record(name, x, y, w, h)
            self._op("image", path=str(path), x=x, y=y, w=w, h=h)
        else:
            self.rect(x, y, w, h, C.PALE, C.MID, True, 1, name=f"{name}_placeholder")
            self.text(f"素材占位\n{label}\n{path.name}", x + .15, y + .15, w - .3, h - .3,
                      14, C.GRAY, False, PP_ALIGN.CENTER, name=f"{name}_label",
                      allow_overlap=True)

    def _preview_font(self, size: float, bold: bool = False):
        """为预览图选择字体；优先使用放在 assets 中的中文字体。"""

        candidates = [
            ASSET_DIR / ("NotoSansCJKsc-Bold.otf" if bold else "NotoSansCJKsc-Regular.otf"),
            ASSET_DIR / "NotoSansCJKsc-Regular.otf",
            Path("/usr/share/fonts/truetype/noto/NotoSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        ]
        for path in candidates:
            if path.exists():
                return ImageFont.truetype(str(path), max(10, int(size * 1.45)))
        return ImageFont.load_default()

    @staticmethod
    def _wrap_for_pixels(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> str:
        """按像素宽度换行，兼容中英文混排。"""

        lines: list[str] = []
        for raw in text.split("\n"):
            if not raw:
                lines.append("")
                continue
            current = ""
            for char in raw:
                trial = current + char
                if current and draw.textlength(trial, font=font) > max_width:
                    lines.append(current.rstrip())
                    current = char.lstrip()
                else:
                    current = trial
            lines.append(current)
        return "\n".join(lines)

    def render_previews(self) -> list[str]:
        """从同一组页面对象生成 PNG，供缺少 Office 渲染器时检查布局。"""

        PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
        output: list[str] = []
        for slide_no, ops in enumerate(self.preview_ops, start=1):
            canvas = Image.new("RGB", (1600, 900), "white")
            draw = ImageDraw.Draw(canvas)
            sx, sy = 1600 / SLIDE_W, 900 / SLIDE_H
            for op in ops:
                kind = op["kind"]
                if kind in {"rect", "roundrect"}:
                    box = [int(op["x"]*sx), int(op["y"]*sy),
                           int((op["x"]+op["w"])*sx), int((op["y"]+op["h"])*sy)]
                    kwargs = {"fill": "#" + op["fill"], "outline": "#" + op["line"],
                              "width": max(1, int(op.get("line_width", 1)*1.5))}
                    if kind == "roundrect": draw.rounded_rectangle(box, radius=14, **kwargs)
                    else: draw.rectangle(box, **kwargs)
                elif kind == "hex":
                    x, y, w, h = op["x"]*sx, op["y"]*sy, op["w"]*sx, op["h"]*sy
                    draw.polygon(hexagon_points(x, y, w, h), fill="#"+op["fill"], outline="#"+op["line"])
                    font = self._preview_font(op["size"], True)
                    text = self._wrap_for_pixels(draw, op["text"], font, int(w*.72))
                    bbox = draw.multiline_textbbox((0, 0), text, font=font, align="center", spacing=2)
                    draw.multiline_text((x+w/2, y+h/2-(bbox[3]-bbox[1])/2), text, font=font,
                                        fill="#"+op["color"], anchor="ma", align="center", spacing=2)
                elif kind == "line":
                    draw.line([op["x1"]*sx, op["y1"]*sy, op["x2"]*sx, op["y2"]*sy],
                              fill="#"+op["color"], width=max(1, int(op["width"]*2)))
                elif kind == "arrow":
                    x, y, w, h = op["x"]*sx, op["y"]*sy, op["w"]*sx, op["h"]*sy
                    if op["direction"] == "right":
                        pts = [(x, y+h*.22), (x+w*.65, y+h*.22), (x+w*.65, y),
                               (x+w, y+h*.5), (x+w*.65, y+h), (x+w*.65, y+h*.78), (x, y+h*.78)]
                    elif op["direction"] == "left":
                        pts = [(x+w, y+h*.22), (x+w*.35, y+h*.22), (x+w*.35, y),
                               (x, y+h*.5), (x+w*.35, y+h), (x+w*.35, y+h*.78), (x+w, y+h*.78)]
                    elif op["direction"] == "down":
                        pts = [(x+w*.22, y), (x+w*.78, y), (x+w*.78, y+h*.65),
                               (x+w, y+h*.65), (x+w*.5, y+h), (x, y+h*.65), (x+w*.22, y+h*.65)]
                    else:
                        pts = [(x+w*.22, y+h), (x+w*.78, y+h), (x+w*.78, y+h*.35),
                               (x+w, y+h*.35), (x+w*.5, y), (x, y+h*.35), (x+w*.22, y+h*.35)]
                    draw.polygon(pts, fill="#"+op["color"])
                elif kind == "text":
                    x, y, w, h = op["x"]*sx, op["y"]*sy, op["w"]*sx, op["h"]*sy
                    font = self._preview_font(op["size"], op["bold"])
                    wrapped = self._wrap_for_pixels(draw, op["text"], font, max(20, int(w-8)))
                    align = "center" if "CENTER" in op["align"] else ("right" if "RIGHT" in op["align"] else "left")
                    anchor_x = x+w/2 if align == "center" else (x+w-4 if align == "right" else x+4)
                    anchor = "ma" if align == "center" else ("ra" if align == "right" else "la")
                    bbox = draw.multiline_textbbox((0, 0), wrapped, font=font, spacing=3, align=align)
                    ty = y + max(2, (h-(bbox[3]-bbox[1]))/2)
                    draw.multiline_text((anchor_x, ty), wrapped, font=font, fill="#"+op["color"],
                                        spacing=3, align=align, anchor=anchor)
                elif kind == "table":
                    rows, cols = len(op["data"]), len(op["data"][0])
                    x, y, w, h = op["x"]*sx, op["y"]*sy, op["w"]*sx, op["h"]*sy
                    cw, rh = w/cols, h/rows
                    for rr, row in enumerate(op["data"]):
                        for cc, value in enumerate(row):
                            box = [x+cc*cw, y+rr*rh, x+(cc+1)*cw, y+(rr+1)*rh]
                            fill = "#"+C.PALE_BLUE if op["header"] and rr == 0 else "white"
                            draw.rectangle(box, fill=fill, outline="#"+C.LIGHT, width=1)
                            font = self._preview_font(op["font_size"], rr == 0)
                            wrapped = self._wrap_for_pixels(draw, str(value), font, int(cw-14))
                            draw.multiline_text((box[0]+7, box[1]+7), wrapped, font=font,
                                                fill="#"+C.INK, spacing=2)
                elif kind == "image":
                    try:
                        img = Image.open(op["path"]).convert("RGB")
                        target = (int(op["w"]*sx), int(op["h"]*sy))
                        img.thumbnail(target)
                        px = int(op["x"]*sx + (target[0]-img.width)/2)
                        py = int(op["y"]*sy + (target[1]-img.height)/2)
                        canvas.paste(img, (px, py))
                    except OSError:
                        pass
            path = PREVIEW_DIR / f"slide_{slide_no:02d}.png"
            canvas.save(path)
            output.append(str(path))
        return output

    def quality_report(self) -> dict:
        """检查越界、明显重叠和过小元素；返回可机器读取的报告。"""

        issues: list[dict] = []
        for rec in self.records:
            if rec.left < 0 or rec.top < 0 or rec.right > SLIDE_W + .001 or rec.bottom > SLIDE_H + .001:
                issues.append({"slide": rec.slide_no, "type": "out_of_bounds", "shape": rec.name,
                               "box": [rec.left, rec.top, rec.width, rec.height]})
            if rec.width < .01 or rec.height < .01:
                issues.append({"slide": rec.slide_no, "type": "too_small", "shape": rec.name})

        # 只对两个都未标记为可重叠的元素检查明显相交，忽略线条、文字覆盖卡片等预期关系。
        for slide_no in range(1, len(self.prs.slides) + 1):
            items = [r for r in self.records if r.slide_no == slide_no and not r.allow_overlap]
            for i, first in enumerate(items):
                for second in items[i + 1:]:
                    ix = max(0.0, min(first.right, second.right) - max(first.left, second.left))
                    iy = max(0.0, min(first.bottom, second.bottom) - max(first.top, second.top))
                    area = ix * iy
                    if area > .04:
                        issues.append({"slide": slide_no, "type": "overlap",
                                       "shapes": [first.name, second.name],
                                       "intersection_area": round(area, 3)})

        # 使用与预览图相同的字体度量估算文本是否会被文本框裁切。
        probe = ImageDraw.Draw(Image.new("RGB", (10, 10)))
        for slide_no, ops in enumerate(self.preview_ops, start=1):
            sx, sy = 1600 / SLIDE_W, 900 / SLIDE_H
            for op in ops:
                if op["kind"] != "text":
                    continue
                font = self._preview_font(op["size"], op["bold"])
                wrapped = self._wrap_for_pixels(probe, op["text"], font, int(op["w"] * sx - 8))
                bbox = probe.multiline_textbbox((0, 0), wrapped, font=font, spacing=3)
                if bbox[3] - bbox[1] > op["h"] * sy - 4:
                    issues.append({"slide": slide_no, "type": "possible_text_overflow",
                                   "text": op["text"][:80], "height_px": bbox[3] - bbox[1],
                                   "box_height_px": round(op["h"] * sy, 1)})

        return {
            "slide_count": len(self.prs.slides),
            "shape_count": len(self.records),
            "issue_count": len(issues),
            "issues": issues,
        }

    def save_notes(self) -> None:
        """把不投屏的口述内容写入独立 Markdown 文件。"""

        lines = ["# Workshop 讲者备注", "", "此文件不会写入 PPT 画布。", ""]
        for idx, meta in enumerate(self.meta, start=1):
            lines.extend([
                f"## {idx}. {meta.title}", "",
                f"- slide_id：`{meta.slide_id}`",
                f"- Stage：{meta.stage}",
                f"- 建议时间：{meta.minutes}",
                f"- 核心 message：{meta.message}",
            ])
            lines.extend(f"- 口述：{item}" for item in meta.notes)
            if meta.transition:
                lines.append(f"- 转场：{meta.transition}")
            lines.append("")
        NOTES_PATH.write_text("\n".join(lines), encoding="utf-8")

    def render_with_libreoffice(self) -> dict:
        """环境具备 LibreOffice 时生成真实 PPT 渲染；否则明确记录降级。"""

        office = shutil.which("libreoffice") or shutil.which("soffice")
        pdftoppm = shutil.which("pdftoppm")
        if not office or not pdftoppm:
            return {"rendered": False, "reason": "未找到 LibreOffice/soffice 或 pdftoppm"}
        RENDER_DIR.mkdir(parents=True, exist_ok=True)
        proc = subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir",
                               str(RENDER_DIR), str(PPTX_PATH)], capture_output=True,
                              text=True, timeout=180, check=False)
        if proc.returncode != 0:
            return {"rendered": False, "reason": proc.stderr.strip()}
        pdf = RENDER_DIR / f"{PPTX_PATH.stem}.pdf"
        proc = subprocess.run([pdftoppm, "-png", "-r", "120", str(pdf),
                               str(RENDER_DIR / "slide")], capture_output=True,
                              text=True, timeout=180, check=False)
        images = sorted(str(p) for p in RENDER_DIR.glob("slide-*.png"))
        return {"rendered": proc.returncode == 0, "pdf": str(pdf), "images": images,
                "reason": proc.stderr.strip() if proc.returncode else ""}

    def finish(self) -> dict:
        """保存 PPT、备注、预览与质量报告。"""

        OUT_DIR.mkdir(parents=True, exist_ok=True)
        ASSET_DIR.mkdir(parents=True, exist_ok=True)
        self.prs.save(PPTX_PATH)
        self.save_notes()
        previews = self.render_previews()
        report = self.quality_report()
        report["pptx_path"] = str(PPTX_PATH)
        report["speaker_notes_path"] = str(NOTES_PATH)
        report["preview_images"] = previews
        report["office_render"] = self.render_with_libreoffice()
        REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return report


def meta(slide_id: str, title: str, stage: str, minutes: str, message: str,
         notes: Sequence[str], transition: str = "") -> SlideMeta:
    """用较短语法构造页面元数据。"""

    return SlideMeta(slide_id, title, stage, minutes, message, list(notes), transition)


def add_bullets(b: DeckBuilder, items: Sequence[str], x: float, y: float, w: float,
                line_h: float = .52, size: float = 19, color: str = C.INK,
                name: str = "bullets") -> None:
    """将短要点排成留白充足的列表。"""

    for idx, item in enumerate(items):
        b.text("•", x, y + idx*line_h, .24, .32, size, C.BLUE, True,
               name=f"{name}_{idx}_dot", allow_overlap=True)
        b.text(item, x + .28, y + idx*line_h, w - .28, .38, size, color,
               name=f"{name}_{idx}_text", allow_overlap=True)


def add_code_card(b: DeckBuilder, code: str, x: float, y: float, w: float, h: float,
                  title: str = "CODE", name: str = "code") -> None:
    """添加深色、原生可编辑的代码卡片。"""

    b.rect(x, y, w, h, C.NAVY, C.NAVY, True, name=f"{name}_bg")
    chip_w = min(w - .36, max(.82, len(title) * .13))
    b.chip(title, x + .18, y + .16, chip_w, C.BLUE, C.WHITE, name=f"{name}_chip")
    b.text(code, x + .24, y + .58, w - .48, h - .78, 14.5, C.WHITE,
           font=FONT_MONO, valign=MSO_ANCHOR.TOP, name=f"{name}_text",
           allow_overlap=True)


def slide_01(b: DeckBuilder) -> None:
    """封面：用一条演化路径建立整场叙事。"""

    b.new_slide(meta("cover", "从一次 API 调用到会自己修订 PPT 的 Agent", "OPEN", "1 分钟",
                     "关注控制权和反馈闭环，而不是 PPT 美化。",
                     ["整场只使用一个市场调研任务。", "持续追问：下一步由程序还是由模型决定？"],
                     "先看系统究竟要完成什么任务。"), show_progress=False)
    b.text("PHD WORKSHOP · AGENT FROM SCRATCH", .72, .55, 5.5, .28, 12, C.BLUE, True,
           name="eyebrow")
    b.text("从一次 API 调用到\n会自己修订 PPT 的 Agent", .72, 1.18, 8.9, 1.55,
           34, C.NAVY, True, valign=MSO_ANCHOR.TOP, name="cover_title")
    b.text("同一个任务，八个阶段：S0 → S7", .76, 2.88, 6.0, .42,
           19, C.GRAY, name="subtitle")
    labels = ["文本草稿", "JSON", "PPT", "联网调研", "自主补搜", "计划调整", "检查修订"]
    start, gap, w = .78, .18, 1.55
    for i, label in enumerate(labels):
        x = start + i*(w+gap)
        fill = C.PALE_BLUE if i in {0, 6} else C.PALE
        line = C.NAVY if i in {0, 6} else C.LIGHT
        b.box(label, x, 4.05, w, .74, fill, line, 15, C.NAVY if i in {0, 6} else C.INK,
              name=f"evolution_{i}")
        if i < len(labels)-1:
            b.arrow(x+w+.03, 4.26, .13, .28, C.MID, name=f"evolution_arrow_{i}")
    b.text("Structured Output · Tools · Workflow · Agent Loop · State · Planning · Reflection",
           .75, 6.55, 11.85, .3, 12, C.GRAY, align=PP_ALIGN.CENTER, name="keywords")


def slide_02(b: DeckBuilder) -> None:
    """统一任务：展示输入、处理与五页交付。"""

    b.new_slide(meta("task", "统一任务：调研一个真实且有口径陷阱的市场", "TASK", "2 分钟",
                     "任务难点是数据可比性，而不是生成五页文字。",
                     ["MAU、DAU、下载量和访问量不能直接混用。", "每个关键数字都必须能够回到已读来源。"],
                     "先给出一个足够简单但实用的 Agent 定义。"), show_progress=False)
    b.rect(.72, 1.42, 4.15, 4.72, C.PALE_BLUE, C.LIGHT, True, name="task_card")
    b.chip("TASK", .98, 1.68, .72, C.NAVY, C.WHITE, name="task_label")
    b.text("截至指定日期，制作 5 页中文 PPT，分析中国大陆消费端通用 AI 助手的竞争格局及最近 12 个月变化。",
           .98, 2.15, 3.63, 1.28, 21, C.NAVY, True, valign=MSO_ANCHOR.TOP,
           name="task_statement", allow_overlap=True)
    add_bullets(b, ["5–6 款独立移动 App", "用户规模优先使用 MAU", "至少 3 个可比时间点", "关键数字可追溯"],
                .98, 3.62, 3.55, .48, 16, name="task_constraints")
    for i, label in enumerate(["调研", "核对", "制作"]):
        y = 2.15 + i*1.12
        b.box(label, 5.28, y, 1.25, .55, C.WHITE, C.BLUE, 16, C.NAVY, name=f"process_{i}")
        if i < 2: b.arrow(5.79, y+.67, .24, .32, C.BLUE, "down", f"process_arrow_{i}")
    outputs = ["范围与结论", "最新格局", "12 个月变化", "产品优劣势", "趋势与局限"]
    b.text("5-PAGE OUTPUT", 7.05, 1.48, 2.3, .3, 12, C.BLUE, True, name="output_label")
    for i, item in enumerate(outputs):
        b.box(f"{i+1:02d}   {item}", 7.05, 1.9+i*.77, 5.32, .57, C.WHITE, C.LIGHT,
              17, C.INK, name=f"output_{i}")
    b.box("MAU ≠ DAU ≠ 下载量 ≠ 访问量", 4.1, 6.39, 5.2, .43,
          C.PALE_RED, C.PALE_RED, 16, C.RED, name="metric_warning")


def slide_03(b: DeckBuilder) -> None:
    """最小定义：Loop 环绕 Model、Tools 与 State。"""

    b.new_slide(meta("agent_definition", "一个用于教学的最小定义", "MODEL", "2 分钟",
                     "Agent 的关键是外部反馈进入持续控制循环。",
                     ["这是教学上的简化表达，不是唯一学术定义。", "复杂输出或自动生成 PPT 都不能单独证明它是 Agent。"],
                     "接下来把这些能力拆成八个可观察阶段。"), show_progress=False)
    b.text("Agent ≈ Model + Tools + State + Loop", 1.0, 1.45, 11.3, .72,
           30, C.NAVY, True, PP_ALIGN.CENTER, name="formula")
    components = [
        ("MODEL", "理解目标\n生成内容\n选择动作", C.PALE_BLUE, C.NAVY),
        ("TOOLS", "接触网络\n文件与程序", C.PALE_ORANGE, C.ORANGE),
        ("STATE", "保存任务事实\n记录已发生什么", C.PALE_GREEN, C.GREEN),
    ]
    for i, (title, body, fill, line) in enumerate(components):
        x = 1.48 + i*3.63
        b.rect(x, 2.62, 3.0, 2.18, fill, line, True, 1.4, name=f"component_{i}")
        b.text(title, x+.22, 2.9, 2.56, .35, 15, line, True, PP_ALIGN.CENTER,
               name=f"component_title_{i}", allow_overlap=True)
        b.text(body, x+.25, 3.42, 2.5, .9, 18, C.INK, False, PP_ALIGN.CENTER,
               name=f"component_body_{i}", allow_overlap=True)
    b.line(1.18, 2.35, 12.0, 2.35, C.BLUE, 2.2, name="loop_top")
    b.line(1.18, 5.1, 12.0, 5.1, C.BLUE, 2.2, name="loop_bottom")
    b.arrow(11.72, 2.23, .35, .24, C.BLUE, name="loop_forward")
    b.arrow(1.05, 4.98, .35, .24, C.BLUE, name="loop_back")
    b.chip("LOOP：Observation 决定 Next Action", 4.45, 4.93, 4.45, C.BLUE, C.WHITE, name="loop_label")
    b.box("复杂输出 ≠ Agent    ·    自动生成 PPT ≠ Agent", 2.6, 5.75, 8.15, .55,
          C.PALE_RED, C.PALE_RED, 17, C.RED, name="not_agent")


def slide_04(b: DeckBuilder) -> None:
    """路线图：两行八级阶梯避免横向拥挤。"""

    b.new_slide(meta("roadmap", "S0–S7：每个阶段只增加一个核心机制", "S0", "2 分钟",
                     "一次只改变一个关键机制，才能看清控制权如何迁移。",
                     ["S1 只生成 JSON；PPT 构建从 S2 才出现。", "State 与 Planning 分开讲：先保存事实，再调整未来动作。"],
                     "从最小的 S0 开始。"))
    stages = [
        ("S0", "回答", "draft.md"), ("S1", "结构化", "deck.json"),
        ("S2", "执行", "tool events"), ("S3", "编排", "workflow trace"),
        ("S4", "自主选择", "agent trace"), ("S5", "记状态", "state.json"),
        ("S6", "调计划", "plan.json"), ("S7", "反馈修订", "deck_v2.pptx"),
    ]
    for i, (stage, verb, artifact) in enumerate(stages):
        row, col = divmod(i, 4)
        x, y = .82 + col*3.08, 1.55 + row*2.18
        fill = C.PALE if i < 3 else (C.PALE_BLUE if i < 7 else C.PALE_GREEN)
        line = C.MID if i < 3 else (C.BLUE if i < 7 else C.GREEN)
        b.rect(x, y, 2.62, 1.58, fill, line, True, 1.2, name=f"stage_{i}")
        b.chip(stage, x+.18, y+.16, .62, line, C.WHITE, name=f"stage_chip_{i}")
        b.text(verb, x+.18, y+.58, 2.26, .42, 20, C.NAVY, True,
               name=f"stage_verb_{i}", allow_overlap=True)
        b.text(artifact, x+.18, y+1.08, 2.26, .25, 12, C.GRAY,
               name=f"stage_artifact_{i}", allow_overlap=True)
        if col < 3: b.arrow(x+2.72, y+.66, .22, .28, C.MID, name=f"stage_arrow_{i}")
    b.text("回答 → 结构化 → 执行 → 编排 → 自主选择 → 记状态 → 调计划 → 按反馈修订",
           1.1, 6.17, 11.1, .38, 15, C.GRAY, align=PP_ALIGN.CENTER, name="roadmap_caption")


def slide_05(b: DeckBuilder) -> None:
    """S0：封闭的一次调用。"""

    b.new_slide(meta("s0_api", "S0 · 一次普通 API 调用", "S0", "2 分钟",
                     "S0 只是一次封闭的文本生成。",
                     ["统一入口读取配置后，只调用模型一次。", "没有搜索工具，也不能证明最新数字可靠。"],
                     "程序还无法稳定使用这段自由文本。"))
    add_code_card(b, "draft = model.generate_text(\n    task_prompt(config),\n    instructions\n)",
                  .78, 1.55, 5.05, 3.65, "S0 / PYTHON", "s0_code")
    items = [("config.task", C.PALE), ("task_prompt()", C.PALE_BLUE),
             ("draft.md", C.PALE_GREEN), ("OpenAIModel", C.PALE_BLUE)]
    for i, (label, fill) in enumerate(items):
        x = 6.25 + (i % 2)*3.0
        y = 1.72 + (i // 2)*1.55
        b.box(label, x, y, 2.48, .72, fill, C.LIGHT, 17, C.NAVY,
              name=f"s0_flow_{i}")
        if i == 0: b.arrow(8.78, y+.22, .38, .28, C.BLUE, name="s0_arrow_0")
        if i == 1: b.arrow(9.26, 2.56, .28, .38, C.BLUE, "down", "s0_arrow_1")
        if i == 3: b.arrow(8.78, y+.22, .38, .28, C.BLUE, "left", "s0_arrow_2")
    add_bullets(b, ["一次请求", "没有工具", "输出 draft.md"], 6.43, 5.05, 5.1,
                .46, 18, name="s0_features")
    b.box("普通 LLM 应用 ≠ Agent", 2.0, 5.63, 3.3, .52, C.PALE_RED, C.PALE_RED,
          17, C.RED, name="s0_conclusion")


def slide_06(b: DeckBuilder) -> None:
    """S0 限制：真实草稿截图与程序解析问号。"""

    b.new_slide(meta("s0_limit", "S0 能写，但不能可靠执行", "S0", "1.5 分钟",
                     "自由文本不是稳定的软件接口。",
                     ["人能读懂草稿，但程序无法依赖漂移的段落格式。", "下一步先解决机器可读性，而不是增加自治。"],
                     "S1 只解决结构问题。"))
    b.placeholder("S0 draft.md 局部", ASSET_DIR / "s0_draft_excerpt.png",
                  .78, 1.48, 5.45, 4.82, "s0_draft")
    b.text("程序解析器", 6.72, 1.55, 2.1, .36, 14, C.BLUE, True, name="parser_label")
    questions = [("页数？", "格式可能漂移"), ("字段？", "标题与正文难定位"), ("来源？", "数字可能无证据")]
    for i, (q, desc) in enumerate(questions):
        y = 2.05 + i*1.15
        b.box(q, 6.72, y, 1.35, .66, C.PALE_RED, C.RED, 20, C.RED,
              name=f"question_{i}")
        b.text(desc, 8.28, y+.08, 3.65, .48, 17, C.INK, name=f"question_desc_{i}")
    b.box("程序怎样知道标题、页面、来源分别在哪里？", 6.72, 5.58, 5.45, .68,
          C.PALE_BLUE, C.BLUE, 17, C.NAVY, name="parser_question")


def slide_07(b: DeckBuilder) -> None:
    """S1：DeckSpec 字段映射到幻灯片线框。"""

    b.new_slide(meta("s1_contract", "S1 · 用 DeckSpec 把文字变成程序契约", "S1", "2 分钟",
                     "Structured Output 让结果可被程序消费，但仍只是数据。",
                     ["JSON Schema 与本地校验共同保证页数和稳定 ID。", "稳定 ID 使 S7 能对特定页面执行 Patch。"],
                     "有了 JSON，并不等于已经创建 PPT。"))
    code = '{\n  "title": "...",\n  "slides": [{\n    "id": "s1",\n    "title": "...",\n    "bullets": ["..."],\n    "source_ids": [],\n    "speaker_notes": "..."\n  }]\n}'
    add_code_card(b, code, .75, 1.4, 5.1, 4.98, "DECKSPEC / JSON", "deck_json")
    b.rect(7.05, 1.55, 5.18, 4.05, C.WHITE, C.NAVY, True, 1.4, name="slide_wireframe")
    b.rect(7.34, 1.88, 4.58, .55, C.PALE_BLUE, C.LIGHT, name="wire_title", allow_overlap=True)
    b.rect(7.34, 2.72, 3.55, 1.65, C.PALE, C.LIGHT, name="wire_body", allow_overlap=True)
    b.rect(11.12, 2.72, .8, 1.65, C.PALE_GREEN, C.LIGHT, name="wire_id", allow_overlap=True)
    b.rect(7.34, 4.66, 4.58, .38, C.PALE, C.LIGHT, name="wire_sources", allow_overlap=True)
    labels = [("title", 6.1, 2.02, 7.34, 2.15), ("bullets", 6.1, 3.22, 7.34, 3.48),
              ("id", 12.0, 3.2, 11.54, 3.48), ("source_ids", 6.03, 4.73, 7.34, 4.85),
              ("speaker_notes", 8.85, 5.85, 9.55, 5.12)]
    for i, (label, tx, ty, lx, ly) in enumerate(labels):
        b.chip(label, tx, ty, 1.18 if label != "speaker_notes" else 1.7,
               C.PALE_BLUE, C.NAVY, name=f"field_{i}")
        b.line(tx + (.59 if label != "speaker_notes" else .85), ty, lx, ly, C.MID, 1,
               name=f"field_line_{i}")
    b.box("固定字段 · 恰好 5 页 · ID = s1–s5", 6.85, 6.0, 5.38, .46,
          C.PALE_GREEN, C.PALE_GREEN, 15, C.GREEN, name="deck_contract")


def slide_08(b: DeckBuilder) -> None:
    """S1/S2 边界：数据与外部执行明确断开。"""

    b.new_slide(meta("s1_boundary", "关键边界：JSON 不是 PPT", "S1", "1.5 分钟",
                     "结构化输出与工具执行是两件不同的事。",
                     ["S1 的验收条件是运行目录中没有 PPTX。", "真正创建文件需要 Runtime 执行 python-pptx。"],
                     "S2 开始让模型通过工具影响外部环境。"))
    b.text("Structured Output  ≠  Tool Calling", 2.0, 1.38, 9.3, .62,
           29, C.NAVY, True, PP_ALIGN.CENTER, name="boundary_statement")
    b.chip("S1 · DATA", .88, 2.35, 1.15, C.MID, C.WHITE, name="s1_label")
    b.box("Model", 2.42, 2.25, 1.45, .72, C.PALE_BLUE, C.BLUE, 18, C.NAVY, name="s1_model")
    b.arrow(4.02, 2.47, .58, .25, C.MID, name="s1_json_arrow")
    b.box("deck.json", 4.78, 2.25, 1.72, .72, C.PALE_GREEN, C.GREEN, 17, C.GREEN, name="s1_json")
    b.line(6.84, 2.1, 6.84, 3.15, C.RED, 2, True, name="boundary_break")
    b.text("STOP", 6.51, 3.2, .68, .28, 12, C.RED, True, PP_ALIGN.CENTER, name="stop_label")
    b.box(".pptx", 7.25, 2.25, 1.45, .72, C.PALE, C.LIGHT, 17, C.MID, name="s1_no_ppt")
    b.chip("S2 · ACTION", .88, 4.18, 1.42, C.NAVY, C.WHITE, name="s2_label")
    nodes = [("Model", C.PALE_BLUE, C.BLUE), ("ToolCall", C.PALE_BLUE, C.BLUE),
             ("Runtime", C.PALE, C.MID), ("create_ppt()", C.PALE_ORANGE, C.ORANGE),
             (".pptx", C.PALE_GREEN, C.GREEN)]
    for i, (label, fill, line) in enumerate(nodes):
        x = 2.42 + i*2.0
        if label == "create_ppt()": b.hex(label, x, 4.05, 1.55, .78, fill, line, 14, f"s2_node_{i}")
        else: b.box(label, x, 4.08, 1.55, .72, fill, line, 16, C.NAVY, name=f"s2_node_{i}")
        if i < len(nodes)-1: b.arrow(x+1.62, 4.3, .28, .24, C.BLUE, name=f"s2_arrow_{i}")
    b.box("S1 到 deck.json 为止", 4.63, 5.65, 4.05, .55, C.PALE_RED, C.PALE_RED,
          17, C.RED, name="s1_end")


def slide_09(b: DeckBuilder) -> None:
    """S2 Tool Calling 时序：突出 Runtime 与 call_id。"""

    b.new_slide(meta("s2_protocol", "Tool Calling：模型提出动作，Runtime 执行动作", "S2", "2.5 分钟",
                     "Tool Calling 是请求、执行与回传协议。",
                     ["模型只提出希望调用的工具和参数。", "Runtime 负责权限、参数校验和真实执行。"],
                     "接下来区分四个常被混用的概念。"))
    lanes = [("USER", 1.0), ("MODEL", 4.0), ("RUNTIME", 7.0), ("TOOL", 10.0)]
    for label, x in lanes:
        line = C.ORANGE if label == "TOOL" else C.BLUE
        b.chip(label, x, 1.42, 1.55, line, C.WHITE, name=f"lane_{label}")
        b.line(x+.78, 1.86, x+.78, 5.86, C.LIGHT, 1.2, True, name=f"lane_line_{label}")
    events = [
        (1.78, 4.78, 2.08, "task", C.BLUE),
        (4.78, 7.78, 2.82, "ToolCall · call_id=42", C.BLUE),
        (7.78, 10.78, 3.55, "execute", C.ORANGE),
        (10.78, 7.78, 4.28, "ToolResult · call_id=42", C.GREEN),
        (7.78, 4.78, 5.02, "result 回传模型", C.GREEN),
    ]
    for i, (x1, x2, y, label, color) in enumerate(events):
        b.line(x1, y, x2, y, color, 2, name=f"event_line_{i}")
        b.arrow(x2-.18 if x2 > x1 else x2, y-.11, .2, .22, color,
                "right" if x2 > x1 else "left", f"event_arrow_{i}")
        b.text(label, min(x1, x2)+.12, y-.42, abs(x2-x1)-.24, .28, 13.5,
               color, True, PP_ALIGN.CENTER, name=f"event_label_{i}")
    b.box("模型不会直接运行 Python，也不会直接访问网络", 2.65, 6.12, 8.0, .52,
          C.PALE_RED, C.PALE_RED, 17, C.RED, name="runtime_boundary")


def slide_10(b: DeckBuilder) -> None:
    """概念边界：同一个工具可被不同控制器调用。"""

    b.new_slide(meta("concepts", "Tool、Tool Calling、Workflow、Agent 的区别", "S2", "2 分钟",
                     "工具相同不代表控制方式相同。",
                     ["Tool 是能力，Tool Calling 是协议。", "Workflow 和 Agent 的核心差异是下一步控制权。"],
                     "S2 用两个有边界的回合把协议跑一遍。"))
    definitions = [("Tool", "系统能够执行什么？"), ("Tool Calling", "模型怎样请求执行？"),
                   ("Workflow", "代码预先规定什么顺序？"),
                   ("Agent Loop", "谁根据 Observation 决定下一步？")]
    for i, (term, question) in enumerate(definitions):
        y = 1.44 + i*1.08
        b.chip(term, .78, y+.08, 1.55, C.NAVY if i == 3 else C.BLUE, C.WHITE,
               name=f"concept_{i}")
        b.text(question, 2.58, y, 4.25, .48, 18, C.INK, bold=i == 3,
               name=f"concept_question_{i}")
    b.box("Python\nWorkflow", 7.45, 1.65, 1.72, 1.0, C.PALE_BLUE, C.BLUE, 17, C.NAVY,
          name="workflow_controller")
    b.box("Model\nAction", 10.25, 1.65, 1.72, 1.0, C.PALE_BLUE, C.BLUE, 17, C.NAVY,
          name="agent_controller")
    b.arrow(8.2, 2.83, .45, .45, C.BLUE, "down", "workflow_to_tool")
    b.arrow(10.98, 2.83, .45, .45, C.BLUE, "down", "agent_to_tool")
    b.hex("search_web", 8.62, 3.54, 2.2, 1.05, C.PALE_ORANGE, C.ORANGE, 17, "shared_tool")
    b.line(8.43, 3.27, 9.18, 3.54, C.BLUE, 2, name="workflow_tool_line")
    b.line(11.2, 3.27, 10.28, 3.54, C.BLUE, 2, name="agent_tool_line")
    b.box("同一个工具，可由两种控制方式调用", 7.35, 5.18, 4.82, .6,
          C.PALE_GREEN, C.PALE_GREEN, 17, C.GREEN, name="shared_tool_message")


def slide_11(b: DeckBuilder) -> None:
    """S2 两个固定回合：使用真实日志素材或明确占位。"""

    b.new_slide(meta("s2_trace", "S2 · 两个有边界的教学回合", "S2", "2.5 分钟",
                     "S2 有 Tool Calling，但没有自主循环。",
                     ["第一个回合只开放 search_web，第二个只开放 create_ppt。", "搜索摘要只是 candidate_only，不能当作已核实证据。"],
                     "下一步将工具串成固定 Workflow。"))
    for key, y, color, title, flow in [
        ("a", 1.45, C.BLUE, "search_web", "Query → candidates → summary"),
        ("b", 3.1, C.ORANGE, "create_ppt", "DeckSpec → PPTX → summary")]:
        b.rect(.75, y, 3.0, 1.38, C.PALE_BLUE if key == "a" else C.PALE_ORANGE,
               color, True, name=f"round_{key}")
        b.chip(f"回合 {key.upper()}", .98, y+.21, .88, color, C.WHITE, name=f"round_{key}_chip")
        b.text(title, 1.0, y+.63, 2.5, .3, 18, C.NAVY, True,
               name=f"round_{key}_title", allow_overlap=True)
        b.text(flow, 1.0, y+.98, 2.5, .24, 12, C.GRAY,
               name=f"round_{key}_flow", allow_overlap=True)
    b.placeholder("S2 events.jsonl", ASSET_DIR / "s2_events.png", 4.15, 1.45, 4.55, 3.95, "s2_events")
    b.placeholder("S2 PPT 单页预览", ASSET_DIR / "s2_deck_preview.png", 9.05, 1.45, 3.45, 3.95, "s2_preview")
    b.box("宿主决定回合起止 · 没有自主循环", 2.35, 5.82, 8.65, .6,
          C.PALE_RED, C.PALE_RED, 18, C.RED, name="s2_not_agent")


def slide_12(b: DeckBuilder) -> None:
    """S3 固定 Workflow：所有箭头保持单向。"""

    b.new_slide(meta("s3_workflow", "S3 · 下一步由 Python 预先决定", "S3", "2.5 分钟",
                     "固定 Workflow 能完成任务，但整体路径由代码控制。",
                     ["模型参与局部选择、抽取和写作。", "路径稳定时，固定 Workflow 往往更可靠、更便宜。"],
                     "真实调研还需要把网页变成可比较证据。"))
    steps = [("搜索排名", "TOOL"), ("选择产品", "MODEL"), ("搜索 MAU", "TOOL"),
             ("读取来源", "TOOL"), ("提取 Evidence", "MODEL"), ("检查口径", "CONTRACT"),
             ("生成 DeckSpec", "MODEL"), ("create_ppt", "TOOL")]
    for i, (label, owner) in enumerate(steps):
        row, col = divmod(i, 4); x, y = .78 + col*3.08, 1.62 + row*1.75
        fill = C.PALE_ORANGE if owner == "TOOL" else (C.PALE_BLUE if owner == "MODEL" else C.PALE_GREEN)
        line = C.ORANGE if owner == "TOOL" else (C.BLUE if owner == "MODEL" else C.GREEN)
        b.box(label, x, y, 2.58, .82, fill, line, 17, C.NAVY, name=f"workflow_step_{i}")
        b.chip(owner, x+.76, y+.95, 1.05, line, C.WHITE, name=f"workflow_owner_{i}")
        if col < 3: b.arrow(x+2.68, y+.27, .24, .28, C.MID, name=f"workflow_arrow_{i}")
        elif row == 0: b.arrow(11.88, y+1.18, .3, .4, C.MID, "down", "workflow_down")
    b.box("稳定 · 可预测 · 易调试", .95, 5.45, 4.12, .55, C.PALE_GREEN, C.PALE_GREEN,
          17, C.GREEN, name="workflow_advantage")
    b.box("资料缺口不会临时改变流程", 7.05, 5.45, 4.9, .55, C.PALE_RED, C.PALE_RED,
          17, C.RED, name="workflow_limit")


def slide_13(b: DeckBuilder) -> None:
    """证据链：候选网页收敛为可比较数据。"""

    b.new_slide(meta("evidence_chain", "搜索摘要不是证据", "S3", "2.5 分钟",
                     "高质量 Agent 依赖数据契约和确定性检查。",
                     ["引用前必须 read_page。", "字段缺失进入 open_questions，而不是猜测。"],
                     "S4 把下一步的控制权交给模型。"))
    stages = [("20 个搜索候选", 5.1, C.PALE_BLUE, C.BLUE), ("12 个已读网页", 4.35, C.PALE_BLUE, C.BLUE),
              ("8 条 Evidence", 3.6, C.PALE_GREEN, C.GREEN), ("2 个同口径组", 2.85, C.PALE_GREEN, C.GREEN),
              ("PPT 图表", 2.1, C.PALE_ORANGE, C.ORANGE)]
    center = 3.55
    for i, (label, width, fill, line) in enumerate(stages):
        y = 1.42 + i*.92
        b.box(label, center-width/2, y, width, .64, fill, line, 16, C.NAVY, name=f"funnel_{i}")
        if i < 4: b.arrow(center-.13, y+.68, .26, .2, C.MID, "down", f"funnel_arrow_{i}")
    b.rect(7.05, 1.45, 5.35, 3.62, C.PALE, C.LIGHT, True, name="evidence_card")
    b.chip("EVIDENCE", 7.32, 1.72, 1.18, C.GREEN, C.WHITE, name="evidence_chip")
    add_bullets(b, ["指标 / 数值 / 单位", "统计期 / 地区 / 终端", "原文片段 / source_id"],
                7.35, 2.35, 4.7, .68, 17, name="evidence_fields")
    b.box("Comparability Report\n按确定性字段分组", 7.55, 4.3, 4.35, .58,
          C.WHITE, C.GREEN, 15, C.GREEN, name="comparability", allow_overlap=True)
    b.box("MAU、DAU、下载量、访问量不能进入同一趋势线", 6.75, 5.55, 5.85, .6,
          C.PALE_RED, C.PALE_RED, 16, C.RED, name="evidence_rule")
    b.text("示意数量", .95, 6.25, 1.2, .22, 10.5, C.GRAY, name="illustrative_label")


def slide_14(b: DeckBuilder) -> None:
    """S4 Agent Loop：四节点闭环与停止出口。"""

    b.new_slide(meta("s4_loop", "S4 · Agent Loop 的最小闭环", "S4", "2.5 分钟",
                     "Observation 必须能改变下一次 Action。",
                     ["Runtime 不硬编码应搜索哪个产品。", "模型结束但未交付 PPT 时只能记为 incomplete。"],
                     "现在比较 Workflow 与 Agent Loop。"))
    b.box("Goal / Context", 1.1, 2.85, 2.2, .82, C.PALE, C.MID, 18, C.NAVY, name="loop_goal")
    b.box("Model chooses\nAction", 4.15, 1.55, 2.3, .95, C.PALE_BLUE, C.BLUE, 18, C.NAVY, name="loop_model")
    b.hex("Runtime executes\nTool", 8.2, 2.75, 2.5, 1.02, C.PALE_ORANGE, C.ORANGE, 16, "loop_tool")
    b.box("Observation\nupdates context", 4.15, 4.25, 2.3, .95, C.PALE_GREEN, C.GREEN, 18, C.GREEN, name="loop_observation")
    b.arrow(3.38, 2.5, .55, .32, C.BLUE, name="loop_a1")
    b.arrow(6.73, 2.16, .66, .32, C.BLUE, name="loop_a2")
    b.arrow(9.2, 3.95, .32, .42, C.GREEN, "down", "loop_a3")
    b.line(8.35, 4.72, 6.55, 4.72, C.GREEN, 2.5, name="loop_return_1")
    b.line(4.0, 4.72, 2.2, 4.72, C.GREEN, 2.5, name="loop_return_2")
    b.arrow(1.95, 3.86, .34, .48, C.GREEN, "up", "loop_return_3")
    b.text("Observation → Next Action", 3.52, 5.55, 4.0, .4, 18, C.GREEN, True,
           PP_ALIGN.CENTER, name="loop_key")
    for i, label in enumerate(["Final response", "PPT delivered", "Limit reached", "Error"]):
        b.chip(label, 8.0 + (i%2)*2.25, 5.35 + (i//2)*.55, 1.95,
               C.PALE_RED if i > 1 else C.PALE_GREEN, C.RED if i > 1 else C.GREEN,
               name=f"exit_{i}")


def slide_15(b: DeckBuilder) -> None:
    """Workflow 与 Agent：相同工具，不同路径。"""

    b.new_slide(meta("workflow_vs_agent", "同一组工具，两种控制方式", "S4", "2 分钟",
                     "更自主不等于一定更好。",
                     ["稳定路径优先使用 Workflow。", "只有路径依赖新信息时，Agent 的动态决策才值得成本。"],
                     "下面看一条真实动态轨迹。"))
    b.text("WORKFLOW", .9, 1.45, 5.25, .34, 14, C.BLUE, True, PP_ALIGN.CENTER, name="workflow_heading")
    b.text("AGENT LOOP", 7.15, 1.45, 5.25, .34, 14, C.BLUE, True, PP_ALIGN.CENTER, name="agent_heading")
    b.rect(.82, 1.92, 5.42, 3.62, C.PALE, C.LIGHT, True, name="workflow_panel", allow_overlap=True)
    b.rect(7.08, 1.92, 5.42, 3.62, C.PALE, C.LIGHT, True, name="agent_panel", allow_overlap=True)
    tools = ["search_web", "read_page", "create_ppt"]
    for i, tool in enumerate(tools):
        x = 1.12 + i*1.72
        b.hex(tool, x, 3.0, 1.45, .75, C.PALE_ORANGE, C.ORANGE, 12.5, f"wf_tool_{i}")
        if i < 2: b.arrow(x+1.48, 3.25, .2, .22, C.MID, name=f"wf_arrow_{i}")
        x2, y2 = 7.38 + i*1.72, 2.55 + (i%2)*1.2
        b.hex(tool, x2, y2, 1.45, .75, C.PALE_ORANGE, C.ORANGE, 12.5, f"agent_tool_{i}")
    b.line(8.85, 2.92, 9.08, 3.55, C.BLUE, 2, name="agent_branch_1")
    b.line(8.85, 2.92, 10.8, 2.92, C.BLUE, 2, name="agent_branch_2")
    b.line(10.8, 3.28, 10.8, 3.75, C.GREEN, 2, name="agent_branch_3")
    b.text("代码预先规定", 2.06, 4.28, 3.0, .34, 17, C.NAVY, True, PP_ALIGN.CENTER,
           name="workflow_control", allow_overlap=True)
    b.text("Observation 决定分支", 8.18, 4.52, 3.2, .34, 17, C.NAVY, True,
           PP_ALIGN.CENTER, name="agent_control", allow_overlap=True)
    b.box("更自主 ≠ 一定更好", 4.35, 5.94, 4.65, .58, C.PALE_RED, C.PALE_RED,
          18, C.RED, name="autonomy_warning")


def slide_16(b: DeckBuilder) -> None:
    """S4 真实轨迹：日志素材与结果页并置。"""

    b.new_slide(meta("s4_trace", "S4 · 发现口径冲突后改变查询", "S4", "3 分钟",
                     "真实事件顺序是动态控制的证据。",
                     ["只展示显式 ToolCall、Observation 和后续动作。", "找不到同口径数据时应缩小范围或明确缺口。"],
                     "循环变长后，需要显式 State。"))
    timeline = [("ACTION", "搜索：AI 助手市场份额", C.PALE_BLUE, C.BLUE),
                ("OBS", "混有下载量、MAU、全球数据", C.PALE_GREEN, C.GREEN),
                ("ACTION", "改搜：中国大陆移动 App MAU", C.PALE_BLUE, C.BLUE),
                ("OBS", "部分产品缺少连续历史数据", C.PALE_GREEN, C.GREEN),
                ("ACTION", "缩小范围 / 标注缺口", C.PALE_BLUE, C.BLUE)]
    for i, (kind, text, fill, line) in enumerate(timeline):
        y = 1.4 + i*.94
        b.chip(kind, .78, y+.16, .82, line, C.WHITE, name=f"trace_kind_{i}")
        b.box(text, 1.82, y, 5.7, .65, fill, line, 16, C.NAVY if kind == "ACTION" else C.GREEN,
              name=f"trace_event_{i}")
        if i < 4: b.arrow(4.48, y+.69, .24, .18, C.MID, "down", f"trace_arrow_{i}")
    b.placeholder("S4 最终趋势页", ASSET_DIR / "s4_result_slide.png", 8.0, 1.42, 4.48, 3.28, "s4_result")
    b.placeholder("S4 events 时间线", ASSET_DIR / "s4_trace.png", 8.0, 4.92, 4.48, 1.28, "s4_trace_asset")
    b.box("后续动作由前一条 Observation 触发", 1.58, 6.28, 5.0, .48,
          C.PALE_GREEN, C.PALE_GREEN, 16, C.GREEN, name="trace_conclusion")


def slide_17(b: DeckBuilder) -> None:
    """S5：把聊天上下文与可信任务状态分开。"""

    b.new_slide(meta("s5_state", "S5 · Context 不等于 State", "S5", "2 分钟",
                     "Context 是本轮输入，State 是可检查的任务事实。",
                     ["State 由 Runtime 根据真实 ToolResult 更新。", "current_deck 只有在 create_ppt 成功后才写入。"],
                     "有了可靠 State，才能讨论计划如何变化。"))
    b.text("CONTEXT", .9, 1.43, 4.25, .34, 14, C.BLUE, True, PP_ALIGN.CENTER, name="context_title")
    b.text("STATE", 8.18, 1.43, 4.25, .34, 14, C.GREEN, True, PP_ALIGN.CENTER, name="state_title")
    b.rect(.85, 1.9, 4.35, 3.95, C.PALE, C.LIGHT, True, name="context_panel")
    messages = ["user: 完成市场调研", "assistant: 调用 search_web", "tool: 返回 10 条结果",
                "assistant: 调用 read_page", "tool: 来源口径不一致"]
    for i, msg in enumerate(messages):
        b.box(msg, 1.15 + (i%2)*.28, 2.18+i*.61, 3.48, .4,
              C.WHITE, C.LIGHT, 13.5, C.GRAY, name=f"message_{i}", allow_overlap=True)
    b.arrow(5.55, 3.16, .78, .48, C.BLUE, name="state_funnel")
    b.text("Runtime\n提取事实", 5.35, 3.72, 1.18, .58, 13, C.BLUE, True,
           PP_ALIGN.CENTER, name="runtime_extract")
    b.rect(7.15, 1.9, 5.35, 3.95, C.PALE_GREEN, C.GREEN, True, name="state_panel")
    b.chip("state.json", 7.48, 2.18, 1.18, C.GREEN, C.WHITE, name="state_file")
    state_fields = ["metric_scope", "source_ids", "open_questions",
                    "completed_actions", "current_deck", "termination_reason"]
    for i, field in enumerate(state_fields):
        col, row = i%2, i//2
        b.box(field, 7.48+col*2.32, 2.83+row*.72, 2.05, .48, C.WHITE, C.LIGHT,
              13.5, C.INK, name=f"state_field_{i}", allow_overlap=True)
    b.box("模型声称“完成” ≠ State 记录“已交付”", 3.05, 6.18, 7.25, .52,
          C.PALE_RED, C.PALE_RED, 16, C.RED, name="state_trust_rule")


def slide_18(b: DeckBuilder) -> None:
    """S6：展示计划在 Observation 前后的差异。"""

    b.new_slide(meta("s6_planning", "S6 · Planning 的价值在 Re-planning", "S6", "2.5 分钟",
                     "计划的价值在于根据新证据改变执行路径。",
                     ["State 描述发生了什么，Plan 描述准备做什么。", "有效 Re-plan 必须改变后续行动，而不只是换一种说法。"],
                     "下一步让最终产物本身接受反馈。"))
    b.text("REVISION 1", .86, 1.42, 4.55, .3, 13, C.BLUE, True, PP_ALIGN.CENTER, name="plan_v1_title")
    b.text("REVISION 2", 7.9, 1.42, 4.55, .3, 13, C.GREEN, True, PP_ALIGN.CENTER, name="plan_v2_title")
    b.rect(.82, 1.82, 4.65, 3.65, C.PALE_BLUE, C.BLUE, True, name="plan_v1_panel")
    b.rect(7.86, 1.82, 4.65, 3.65, C.PALE_GREEN, C.GREEN, True, name="plan_v2_panel")
    v1 = ["选择单一来源", "覆盖 12 个月趋势", "直接绘制一条折线"]
    v2 = ["寻找同口径来源", "无法补齐则缩小范围", "拆图并标注数据缺口"]
    for i, item in enumerate(v1):
        b.box(f"{i+1:02d}  {item}", 1.15, 2.28+i*.88, 3.98, .58, C.WHITE, C.LIGHT,
              15.5, C.INK, name=f"plan_v1_{i}", allow_overlap=True)
    for i, item in enumerate(v2):
        b.box(f"{i+1:02d}  {item}", 8.2, 2.28+i*.88, 3.98, .58,
              C.PALE_ORANGE if i > 0 else C.WHITE, C.ORANGE if i > 0 else C.LIGHT,
              15.5, C.INK, name=f"plan_v2_{i}", allow_overlap=True)
    b.box("OBSERVATION\n中途统计范围改变\n数据不可直接连线", 5.78, 2.5, 1.72, 1.65,
          C.PALE_RED, C.RED, 14, C.RED, name="plan_trigger")
    b.arrow(5.42, 3.16, .3, .3, C.RED, "left", "plan_trigger_left")
    b.arrow(7.52, 3.16, .3, .3, C.RED, name="plan_trigger_right")
    b.box("Re-plan 必须有依据，不能只改措辞", 3.45, 5.98, 6.45, .58,
          C.PALE_RED, C.PALE_RED, 17, C.RED, name="replan_rule")


def slide_19(b: DeckBuilder) -> None:
    """S7：生成、渲染、检查、修订和复查。"""

    b.new_slide(meta("s7_reflection", "S7 · 从生成到检查、修订和复查", "S7", "2 分钟",
                     "Reflection 是外部反馈进入可执行修订循环。",
                     ["确定性规则先发现问题，模型再根据 Issue 决定修改。", "修订后必须重新生成并复查。"],
                     "下面看修改前后是否真的发生变化。"))
    nodes = [
        ("create_ppt", "tool", C.PALE_ORANGE, C.ORANGE),
        ("render_ppt", "tool", C.PALE_ORANGE, C.ORANGE),
        ("check_deck", "check", C.PALE_GREEN, C.GREEN),
        ("ReviewIssue[]", "data", C.PALE_RED, C.RED),
        ("patch_deck", "tool", C.PALE_ORANGE, C.ORANGE),
    ]
    positions = [(1.0, 2.05), (3.35, 2.05), (5.7, 2.05), (8.05, 2.05), (10.4, 2.05)]
    for i, ((label, kind, fill, line), (x, y)) in enumerate(zip(nodes, positions)):
        if kind == "tool": b.hex(label, x, y, 1.75, .82, fill, line, 14, f"review_node_{i}")
        else: b.box(label, x, y, 1.75, .82, fill, line, 14.5, line, name=f"review_node_{i}")
        if i < 4: b.arrow(x+1.83, y+.27, .34, .26, C.BLUE if i < 2 else C.GREEN,
                          name=f"review_arrow_{i}")
    b.line(11.25, 3.12, 11.25, 4.68, C.GREEN, 2.4, name="review_loop_down")
    b.line(11.25, 4.68, 1.86, 4.68, C.GREEN, 2.4, name="review_loop_back")
    b.arrow(1.67, 3.46, .38, .48, C.GREEN, "up", "review_loop_up")
    b.text("create + check again", 4.9, 4.77, 3.4, .34, 17, C.GREEN, True,
           PP_ALIGN.CENTER, name="review_again")
    checks = ["五页结构与文字密度", "数字来源与指标标签", "未知 source_id", "元素越界"]
    for i, label in enumerate(checks):
        b.chip(label, 1.18+i*3.0, 5.7, 2.72, C.PALE_BLUE, C.NAVY, name=f"check_{i}")
    b.text("max_review_rounds", 10.52, 6.24, 1.72, .24, 11.5, C.GRAY,
           False, PP_ALIGN.RIGHT, name="review_limit")


def slide_20(b: DeckBuilder) -> None:
    """S7 前后对比：真实素材缺失时使用明确占位。"""

    b.new_slide(meta("s7_before_after", "ReviewIssue 必须对应可见变化", "S7", "2 分钟",
                     "反馈改变产物并通过复查，才形成闭环。",
                     ["DeckPatch 通过 slide_id 和 issue_ids 保持可追踪。", "找不到可靠来源时宁可删除数字，也不能编造引用。"],
                     "Agent 越自主，越需要明确边界和评价。"))
    b.text("BEFORE", .86, 1.4, 4.75, .3, 13, C.RED, True, PP_ALIGN.CENTER, name="before_title")
    b.text("AFTER", 7.72, 1.4, 4.75, .3, 13, C.GREEN, True, PP_ALIGN.CENTER, name="after_title")
    b.placeholder("S7 修订前页面", ASSET_DIR / "s7_before.png", .82, 1.82, 4.82, 3.52, "s7_before")
    b.placeholder("S7 修订后页面", ASSET_DIR / "s7_after.png", 7.68, 1.82, 4.82, 3.52, "s7_after")
    b.arrow(5.94, 2.85, 1.42, .58, C.GREEN, name="before_after_arrow")
    b.box("ISSUE\n份额标签过度推断\n数字来源缺失", 5.65, 3.65, 2.0, 1.02,
          C.PALE_RED, C.RED, 13.5, C.RED, name="review_issue")
    b.text("市场份额：32%", 1.25, 5.57, 3.85, .3, 16, C.RED, True,
           PP_ALIGN.CENTER, name="before_claim")
    b.text("样本内 MAU 份额：32% · 来源：src_xxx", 8.0, 5.57, 4.2, .3, 15, C.GREEN, True,
           PP_ALIGN.CENTER, name="after_claim")
    b.box("issue_1 → DeckPatch(s2) → deck_v2.pptx → review_v2.json", 2.22, 6.18, 8.9, .5,
          C.PALE_BLUE, C.PALE_BLUE, 15, C.NAVY, name="patch_trace")


def slide_21(b: DeckBuilder) -> None:
    """Agent 的预算、停止条件与评价指标。"""

    b.new_slide(meta("boundaries_eval", "自主行动必须有边界", "EVAL", "1.5 分钟",
                     "Runtime 必须验证交付物并限制资源。",
                     ["模型不能仅靠口头声明完成任务。", "一次精彩轨迹不能证明平均性能提升。"],
                     "最后回到开场的判断标准。"), show_progress=False)
    b.text("LIMITS", .85, 1.44, 5.35, .3, 13, C.RED, True, PP_ALIGN.CENTER, name="limits_title")
    b.text("EVALUATION", 7.15, 1.44, 5.35, .3, 13, C.GREEN, True, PP_ALIGN.CENTER, name="eval_title")
    b.rect(.82, 1.86, 5.42, 3.82, C.PALE_RED, C.RED, True, name="limits_panel")
    limits = ["max_model_calls", "max_tool_calls", "max_agent_steps", "max_elapsed_seconds", "max_review_rounds"]
    for i, label in enumerate(limits):
        b.box(label, 1.18, 2.2+i*.62, 4.7, .42, C.WHITE, C.LIGHT, 14.5, C.INK,
              name=f"limit_{i}", allow_overlap=True)
    b.rect(7.08, 1.86, 5.42, 3.82, C.PALE_GREEN, C.GREEN, True, name="eval_panel")
    evals = ["真实交付 PPT", "数字可追溯、可比较", "Observation 改变动作", "Issue 在修订后消失", "步数、耗时与成本"]
    for i, label in enumerate(evals):
        b.box(label, 7.44, 2.2+i*.62, 4.7, .42, C.WHITE, C.LIGHT, 14.5, C.INK,
              name=f"eval_{i}", allow_overlap=True)
    for i, label in enumerate(["循环", "漂移", "误读", "成本失控"]):
        b.chip(label, 2.0+i*2.4, 6.08, 1.85, C.PALE_RED, C.RED, name=f"risk_{i}")


def slide_22(b: DeckBuilder) -> None:
    """总结：把八个阶段与产物重新连成一条链。"""

    b.new_slide(meta("takeaways", "Takeaways：判断 Agent，看控制权和反馈链", "END", "1.5 分钟",
                     "判断 Agent，要看谁根据反馈决定下一步。",
                     ["确定性检查留给代码，开放式决策交给模型。", "实际系统不必追求最高自治。"],
                     "进入提问或展示运行目录。"), show_progress=False)
    takeaways = [
        ("01", "Structured Output", "让结果可被程序读取"),
        ("02", "Tools", "让模型接触外部世界"),
        ("03", "Workflow", "由代码预先规定路径"),
        ("04", "Agent Loop", "根据 Observation 选择下一步"),
        ("05", "State + Plan", "保存事实并描述未来动作"),
        ("06", "Reflection", "用反馈改变并复查产物"),
    ]
    for i, (num, term, desc) in enumerate(takeaways):
        col, row = i % 2, i // 2
        x, y = .85 + col*6.15, 1.46 + row*1.12
        b.chip(num, x, y+.15, .52, C.NAVY, C.WHITE, name=f"takeaway_num_{i}")
        b.text(term, x+.72, y, 2.05, .38, 17, C.NAVY, True,
               name=f"takeaway_term_{i}")
        b.text(desc, x+.72, y+.4, 4.75, .34, 15.5, C.GRAY,
               name=f"takeaway_desc_{i}")
    artifacts = ["draft", "JSON", "tools", "workflow", "agent", "state", "plan", "deck v2"]
    for i, label in enumerate(artifacts):
        x = .85 + i*1.55
        fill = C.PALE_GREEN if i == 7 else (C.PALE_BLUE if i >= 4 else C.PALE)
        line = C.GREEN if i == 7 else (C.BLUE if i >= 4 else C.LIGHT)
        b.box(f"S{i}\n{label}", x, 5.02, 1.18, .72, fill, line, 13.5, C.NAVY,
              name=f"artifact_{i}")
        if i < 7: b.arrow(x+1.23, 5.25, .22, .24, C.MID, name=f"artifact_arrow_{i}")
    b.box("判断 Agent，不看它做了多少事，看谁决定下一步", 2.0, 6.12, 9.35, .58,
          C.NAVY, C.NAVY, 19, C.WHITE, name="final_message")


def slide_23(b: DeckBuilder) -> None:
    """备用 A：项目目录与职责边界。"""

    b.new_slide(meta("appendix_architecture", "项目目录与职责边界", "APP", "备用",
                     "共享能力与控制流分离。",
                     ["Stage 只定义控制流，Tool 只完成一次动作。", "每次运行保存配置快照和事件流。"]),
                show_progress=False)
    tree = "src/\n├─ run.py / config.yaml\n├─ core/\n├─ stages/\n├─ tools/\n└─ runs/"
    add_code_card(b, tree, .82, 1.48, 4.7, 4.78, "PROJECT TREE", "project_tree")
    roles = [("core/", "配置 · 契约 · 模型 · Runtime", C.BLUE),
             ("stages/", "S0–S7 控制流", C.NAVY),
             ("tools/", "搜索 · PPT · Review", C.ORANGE),
             ("runs/", "隔离的运行产物", C.GREEN)]
    for i, (folder, desc, color) in enumerate(roles):
        y = 1.55 + i*1.05
        b.chip(folder, 6.05, y+.14, 1.2, color, C.WHITE, name=f"folder_{i}")
        b.box(desc, 7.48, y, 4.55, .68, C.WHITE, color, 16, C.INK,
              name=f"folder_desc_{i}")
    b.box("Stage 定义控制流 · Tool 只完成一次动作", 6.02, 5.86, 6.08, .52,
          C.PALE_BLUE, C.PALE_BLUE, 16, C.NAVY, name="architecture_rule")


def slide_24(b: DeckBuilder) -> None:
    """备用 B：四组核心数据契约。"""

    b.new_slide(meta("appendix_contracts", "核心数据契约", "APP", "备用",
                     "数据契约连接概率模型与确定性程序。",
                     ["Schema 约束格式，但不能自动保证事实正确。", "可信度仍需要来源、Evidence 与 Review。"]),
                show_progress=False)
    pairs = [("ToolCall", "ToolResult", "调用"), ("Source", "Evidence", "研究"),
             ("DeckSpec", "PPT", "制作"), ("ReviewIssue", "DeckPatch", "修订")]
    for i, (left, right, role) in enumerate(pairs):
        y = 1.5 + i*1.12
        b.chip(role, .9, y+.19, .78, C.NAVY, C.WHITE, name=f"contract_role_{i}")
        b.box(left, 2.08, y, 3.55, .72, C.PALE_BLUE if i != 3 else C.PALE_RED,
              C.BLUE if i != 3 else C.RED, 18, C.NAVY, name=f"contract_left_{i}")
        b.arrow(5.9, y+.24, .78, .28, C.GREEN, name=f"contract_arrow_{i}")
        b.box(right, 7.0, y, 3.55, .72, C.PALE_GREEN if i != 3 else C.PALE_ORANGE,
              C.GREEN if i != 3 else C.ORANGE, 18, C.NAVY, name=f"contract_right_{i}")
    b.box("Schema 约束格式 · 来源与检查约束可信度", 3.2, 6.05, 6.95, .55,
          C.PALE_GREEN, C.PALE_GREEN, 17, C.GREEN, name="schema_rule")


def slide_25(b: DeckBuilder) -> None:
    """备用 C：配置与零参数运行。"""

    b.new_slide(meta("appendix_config", "唯一配置与零参数运行", "APP", "备用",
                     "配置集中、密钥隔离、入口保持不变。",
                     ["config.yaml 存非敏感业务配置，.env 存 URL 与 Key。", "现场不得展示真实密钥。"]),
                show_progress=False)
    code = "# src/core/.env\nDEEPSEEK_API_KEY=...\nDEEPSEEK_BASE_URL=https://api.deepseek.com\nTHIRD_PARTY_API_KEY=...\nTAVILY_API_KEY=...\n\npython src/run.py"
    add_code_card(b, code, .8, 1.45, 5.2, 4.95, "CONFIG / SHELL", "config_code")
    b.box("config.yaml", 6.55, 1.62, 2.2, .72, C.PALE_BLUE, C.BLUE, 18, C.NAVY, name="config_file")
    b.box("core/.env", 9.5, 1.62, 2.2, .72, C.PALE_RED, C.RED, 18, C.RED, name="env_file")
    b.arrow(7.45, 2.58, .4, .55, C.BLUE, "down", "config_down")
    b.arrow(10.4, 2.58, .4, .55, C.RED, "down", "env_down")
    b.box("run.py", 7.85, 3.35, 2.55, .78, C.NAVY, C.NAVY, 21, C.WHITE, name="run_entry")
    b.text("只改变 run.stage", 7.75, 4.48, 2.8, .32, 16, C.BLUE, True,
           PP_ALIGN.CENTER, name="stage_switch")
    for i in range(8):
        b.chip(f"S{i}", 6.48+i*.72, 5.22, .55, C.PALE_BLUE if i < 7 else C.PALE_GREEN,
               C.NAVY if i < 7 else C.GREEN, name=f"config_stage_{i}")


def slide_26(b: DeckBuilder) -> None:
    """备用 D：明确最小教学 Demo 的范围。"""

    b.new_slide(meta("appendix_scope", "哪些内容没有进入主 Demo", "APP", "备用",
                     "先把单 Agent 的数据与控制边界讲清楚。",
                     ["这些能力并非不重要，只是不应同时进入最小 Demo。", "单 Agent 清晰后，才能解释 Multi-Agent 的额外价值。"]),
                show_progress=False)
    b.text("当前主线", .88, 1.48, 2.0, .32, 14, C.BLUE, True, name="current_track_title")
    for i in range(8):
        x = .88 + i*1.48
        b.box(f"S{i}", x, 2.02, 1.06, .62, C.PALE_BLUE if i < 7 else C.PALE_GREEN,
              C.BLUE if i < 7 else C.GREEN, 17, C.NAVY, name=f"current_stage_{i}")
        if i < 7: b.arrow(x+1.12, 2.22, .2, .22, C.MID, name=f"current_stage_arrow_{i}")
    b.text("后续扩展", .88, 3.34, 2.0, .32, 14, C.GRAY, True, name="future_track_title")
    future = ["Multi-Agent", "Agent Framework", "向量数据库与长期记忆", "图片生成与复杂美化", "生产级重试与权限"]
    for i, label in enumerate(future):
        x = .88 + (i%3)*3.88
        y = 3.92 + (i//3)*1.05
        b.rect(x, y, 3.42, .66, C.PALE, C.MID, True, 1, name=f"future_{i}")
        shape = b.slide.shapes[-1]
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        b.text(label, x+.12, y+.12, 3.18, .4, 15.5, C.GRAY, False, PP_ALIGN.CENTER,
               name=f"future_text_{i}", allow_overlap=True)
    b.box("先把单 Agent 的数据与控制边界讲清楚", 3.2, 6.1, 6.95, .55,
          C.NAVY, C.NAVY, 18, C.WHITE, name="scope_message")


V2_STAGES = [
    {
        "stage": "S0", "arch_title": "S0 · 只有一次普通 API 调用",
        "bridge_title": "S0 · 能写草稿，但程序无法可靠使用",
        "mechanism": "一次 Input → Model → Text",
        "sources": ["stages/s0_api.py", "model.py: generate_text()", "recorder.py: write_text()"],
        "outputs": ["draft.md", "原始请求 / 响应", "events.jsonl"],
        "gaps": ["自由文本格式会漂移", "没有稳定字段交给 PPT 构建器", "没有真实搜索与来源核对"],
        "next": "因为仍缺少机器可读的输出契约，S1 加入 Structured Output。",
        "message": "S0 是一次封闭请求，不具备工具或循环。",
        "notes": ["宿主只决定调用模型一次。", "自由文本对人友好，但不是稳定的软件接口。"],
    },
    {
        "stage": "S1", "arch_title": "S1 · 在模型输出后加入 DeckSpec 契约",
        "bridge_title": "S1 · 能输出合法 JSON，但还不能生成 PPT",
        "mechanism": "Structured Output + 本地校验",
        "sources": ["stages/s1_structured.py", "contracts.py: deck_schema() / validate_deck()", "model.py: generate_structured()"],
        "outputs": ["deck.json", "恰好五页", "稳定 ID：s1–s5"],
        "gaps": ["JSON 不会自行创建文件", "模型不能直接访问网络或运行 Python", "运行目录中不应出现 .pptx"],
        "next": "因为仍缺少影响外部环境的能力，S2 加入 Tools 与 Tool Calling。",
        "message": "S1 改变输出契约，而不是执行能力。",
        "notes": ["Schema 规定字段，本地 validate_deck 再检查一次。", "S1 的验收条件是没有 PPTX。"],
    },
    {
        "stage": "S2", "arch_title": "S2 · 加入 Runtime 和可调用工具",
        "bridge_title": "S2 · 能搜索、能创建 PPT，但还不能独立完成任务",
        "mechanism": "Tool Schema → ToolCall → Execute → ToolResult",
        "sources": ["stages/s2_tools.py", "runtime.py: run_tool_round() / execute_call()", "tools/research.py", "tools/presentation.py"],
        "outputs": ["ToolCall / ToolResult", "sources.json", "示例 deck_v1.pptx"],
        "gaps": ["宿主规定两个回合的起止", "搜索摘要只是候选来源", "没有端到端调研顺序"],
        "next": "因为仍缺少端到端执行顺序，S3 加入固定 Workflow。",
        "message": "模型提出请求，Runtime 校验并执行，Tool 返回真实结果。",
        "notes": ["相同 call_id 连接一次请求与结果。", "有工具不等于已经具备 Agent Loop。"],
    },
    {
        "stage": "S3", "arch_title": "S3 · 用固定 Workflow 编排完整任务",
        "bridge_title": "S3 · 能完成任务，但路线无法根据结果改变",
        "mechanism": "下一步由 Python 预先决定",
        "sources": ["stages/s3_workflow.py", "common.py: run_fixed_workflow()", "research.py: read_page()", "contracts.py: evidence_is_comparable()"],
        "outputs": ["sources / evidence", "comparability.json", "完整五页 PPT"],
        "gaps": ["资料缺口只会被记录", "不能临时改变搜索路线", "异常处理依赖预设分支"],
        "next": "因为仍缺少 Observation 驱动的动态选择，S4 加入 Agent Loop。",
        "message": "S3 能完整交付，但整体路径由 Python 控制。",
        "notes": ["Workflow 不是低级方案，稳定路径往往更适合它。", "此阶段新增的是固定编排，不是新工具。"],
    },
    {
        "stage": "S4", "arch_title": "S4 · 把固定控制流替换为 Agent Loop",
        "bridge_title": "S4 · 能动态补搜，但进度仍藏在消息历史里",
        "mechanism": "Action → Observation → Next Action",
        "sources": ["stages/s4_agent.py", "common.py: run_research_agent()", "runtime.py: run_agent_loop()"],
        "outputs": ["动态 events.jsonl", "自主行动轨迹", "完整 PPT + 结束状态"],
        "gaps": ["消息历史持续增长", "任务事实没有稳定字段", "模型声称完成不等于真实交付"],
        "next": "因为仍缺少可检查的任务事实，S5 加入显式 State。",
        "message": "S4 的关键变化是下一步控制权迁移到模型。",
        "notes": ["真实证据是 Observation 改变后续 Action。", "只展示显式动作，不展示隐藏思维链。"],
    },
    {
        "stage": "S5", "arch_title": "S5 · 在 Agent Loop 旁加入可信 TaskState",
        "bridge_title": "S5 · 能保存可信进度，但还没有未来计划",
        "mechanism": "Runtime 根据真实结果更新显式 State",
        "sources": ["stages/s5_state.py", "contracts.py: TaskState", "context.py: persist_state()", "runtime.py: _record_tool_result()"],
        "outputs": ["S4 全部产物", "持续更新的 state.json", "可追溯结束原因"],
        "gaps": ["没有步骤依赖与完成条件", "没有显式未来路径", "新证据无法产生计划版本"],
        "next": "因为仍缺少可观察、可修订的未来路径，S6 加入 Planning。",
        "message": "State 由 Runtime 根据执行事实维护。",
        "notes": ["current_deck 只有在 create_ppt 成功后更新。", "State 描述发生了什么，不描述未来动作。"],
    },
    {
        "stage": "S6", "arch_title": "S6 · 在 State 之上加入 Planning 与 Re-planning",
        "bridge_title": "S6 · 能根据证据改计划，但还没检查最终产物",
        "mechanism": "Plan → Execute → Observation → update_plan",
        "sources": ["stages/s6_planning.py", "contracts.py: PlanState / plan_schema()", "common.py: create_initial_plan()", "runtime.py: _update_plan_handler()"],
        "outputs": ["plan.json", "revision history", "有依据的 change reason"],
        "gaps": ["PPT 版面尚未反馈给 Agent", "来源与标签需要外部检查", "修改后尚未再次验证"],
        "next": "因为仍缺少产物级反馈闭环，S7 加入 Reflection。",
        "message": "Planning 的价值在于新证据改变执行路径。",
        "notes": ["措辞变化不算 Re-plan，执行路径变化才算。", "旧计划进入 history，不静默覆盖。"],
    },
    {
        "stage": "S7", "arch_title": "S7 · 把最终 PPT 接回反馈循环",
        "bridge_title": "S7 · 反馈已经改变产物，并被再次验证",
        "mechanism": "Generate → Render → Check → Patch → Generate Again",
        "sources": ["stages/s7_reflection.py", "presentation.py: render_ppt() / patch_deck()", "review.py: check_deck()", "contracts.py: apply_deck_patches()"],
        "outputs": ["deck_v1 / deck_v2", "review_v1 / review_v2", "DeckPatch + 可见变化"],
        "gaps": ["生产级权限与故障恢复", "跨任务记忆与系统评测", "Multi-Agent 与框架迁移"],
        "next": "S7 已完成本次 Agent 闭环；再往后进入生产化与规模化问题。",
        "message": "反馈改变产物并通过复查，才形成完整闭环。",
        "notes": ["Reflection 不是自我评价文字。", "无法解决时应保留缺口，不能编造来源。"],
    },
]

# 输出页底部只保留适合投影的一句桥接语；完整因果说明仍写入讲者备注。
V2_NEXT_LABELS = [
    "缺少机器可读契约 → S1：Structured Output",
    "缺少外部执行能力 → S2：Tools / Tool Calling",
    "缺少端到端顺序 → S3：Fixed Workflow",
    "缺少动态选择 → S4：Agent Loop",
    "缺少可检查任务事实 → S5：State",
    "缺少可修订未来路径 → S6：Planning",
    "缺少产物级反馈闭环 → S7：Reflection",
    "本次闭环完成 → 后续进入生产化与规模化",
]

def v2_new_node(b: DeckBuilder, text: str, x: float, y: float, w: float, h: float,
                kind: str, born: int, current: int, name: str,
                size: float | None = None) -> None:
    """按首次出现 Stage 统一绘制新增或继承模块。"""

    color = {"control": C.NAVY, "model": C.BLUE, "runtime": C.BLUE,
             "tool": C.ORANGE, "data": C.GREEN, "issue": C.RED}[kind]
    pale = {"control": C.PALE_BLUE, "model": C.PALE_BLUE, "runtime": C.PALE,
            "tool": C.PALE_ORANGE, "data": C.PALE_GREEN, "issue": C.PALE_RED}[kind]
    is_new = born == current
    fill, line, text_color = (color, color, C.WHITE) if is_new else (pale, color, C.INK)
    if kind == "tool":
        b.hex(text, x, y, w, h, fill, line, size or 12.5, name, text_color=text_color)
    else:
        b.box(text, x, y, w, h, fill, line, size or 13.5, text_color, name=name)
    if is_new:
        b.chip("+ NEW", x + w - .58, y - .18, .62, color, C.WHITE, name=f"{name}_new")


def v2_architecture_page(b: DeckBuilder, spec: dict, stage_index: int) -> None:
    """使用固定六列锚点绘制某个 Stage 的累计系统结构。"""

    b.new_slide(meta(f"s{stage_index}_architecture", spec["arch_title"], spec["stage"],
                     "2.5 分钟" if stage_index in {2, 3, 4, 6, 7} else "2 分钟",
                     spec["message"], spec["notes"], "下一页检查当前输出与关键缺口。"))
    b.text(f"新增机制：{spec['mechanism']}", .75, 1.22, 8.8, .34, 16, C.BLUE, True,
           name="mechanism")
    headings = [("TASK", .72), ("CONTROLLER", 2.35), ("MODEL / CONTRACT", 4.15),
                ("RUNTIME / TOOLS", 6.12), ("STATE / PLAN / REVIEW", 8.35),
                ("ARTIFACTS", 10.78)]
    for label, x in headings:
        b.text(label, x, 1.62, 1.6 if x < 10 else 1.8, .22, 9.5, C.GRAY, True,
               PP_ALIGN.CENTER, name=f"heading_{label}")

    # 这三个节点从 S0 开始一直存在，只更换 Controller 的 Stage 名称。
    v2_new_node(b, "config.yaml\ntask_prompt()", .72, 2.08, 1.42, .82, "control", 0, stage_index, "task")
    v2_new_node(b, f"run.py\n{spec['stage'].lower()} stage", 2.35, 2.08, 1.5, .82,
                "control", stage_index, stage_index, "controller")
    v2_new_node(b, "OpenAIModel", 4.18, 2.08, 1.52, .82, "model", 0, stage_index, "model")
    b.arrow(2.14, 2.35, .18, .24, C.MID, name="task_to_controller")
    b.arrow(3.91, 2.35, .2, .24, C.MID, name="controller_to_model")

    if stage_index >= 1:
        v2_new_node(b, "DeckSpec\nSchema + Validate", 4.18, 3.28, 1.52, .86,
                    "data", 1, stage_index, "contract", 12.0)
        b.arrow(4.78, 2.95, .28, .24, C.GREEN, "down", "model_to_contract")

    if stage_index >= 2:
        v2_new_node(b, "Runtime\nexecute_call", 6.18, 2.08, 1.52, .82,
                    "runtime", 2, stage_index, "runtime")
        v2_new_node(b, "Research\nTools", 6.02, 3.32, .92, .78,
                    "tool", 2, stage_index, "research_tools")
        v2_new_node(b, "PPT\nTools", 7.02, 3.32, .92, .78,
                    "tool", 2, stage_index, "ppt_tools")
        b.arrow(5.76, 2.35, .34, .24, C.BLUE, name="tool_call")
        b.text("ToolCall", 5.62, 2.02, .6, .18, 8.5, C.BLUE, True, PP_ALIGN.CENTER,
               name="tool_call_label")
        b.arrow(6.79, 2.96, .26, .25, C.ORANGE, "down", "runtime_to_tools")

    if stage_index >= 3:
        # S3 的固定顺序在 S4 后保留为浅灰对照，而不是从结构图中消失。
        fill = C.NAVY if stage_index == 3 else C.PALE
        line = C.NAVY if stage_index == 3 else C.MID
        color = C.WHITE if stage_index == 3 else C.GRAY
        b.box("Fixed Workflow\n1 → 2 → … → 8", 2.35, 3.28, 1.5, .86,
              fill, line, 12.5, color, name="fixed_workflow")
        if stage_index == 3:
            b.chip("+ NEW", 3.27, 3.1, .62, C.NAVY, C.WHITE, name="workflow_new")

    if stage_index >= 4:
        v2_new_node(b, "Agent Loop\nAction ↔ Obs.", 2.35, 4.34, 1.5, .86,
                    "control", 4, stage_index, "agent_loop", 12.0)
        # 可观察的 Observation 回环始终指回 Model。
        b.line(7.5, 4.78, 4.95, 4.78, C.GREEN, 2.2, name="observation_line")
        b.arrow(4.72, 4.65, .25, .26, C.GREEN, "left", "observation_arrow")
        b.text("ToolResult / Observation", 5.28, 4.48, 1.9, .2, 9.5, C.GREEN, True,
               PP_ALIGN.CENTER, name="observation_label")

    if stage_index >= 5:
        v2_new_node(b, "TaskState", 8.48, 2.08, 1.55, .72, "data", 5, stage_index, "state")
        b.line(7.75, 2.45, 8.42, 2.45, C.GREEN, 1.8, name="persist_state_line")
        b.text("Persist", 7.82, 2.18, .52, .18, 8.5, C.GREEN, True, name="persist_state_label")
    if stage_index >= 6:
        v2_new_node(b, "PlanState\nrevision history", 8.48, 3.12, 1.55, .82,
                    "data", 6, stage_index, "plan", 11.5)
    if stage_index >= 7:
        v2_new_node(b, "ReviewIssue\nDeckPatch", 8.48, 4.18, 1.55, .82,
                    "issue", 7, stage_index, "review")

    artifact = ["draft.md", "deck.json", "tool events\ndeck_v1.pptx", "evidence\nfull PPT",
                "agent trace\nfull PPT", "state.json\nfull PPT", "plan.json\nfull PPT",
                "deck_v2\nreview_v2"][stage_index]
    artifact_kind = "issue" if stage_index == 7 else "data"
    v2_new_node(b, artifact, 10.78, 2.08, 1.8, .9, artifact_kind,
                stage_index, stage_index, "artifact")
    if stage_index == 0:
        b.arrow(5.76, 2.35, 4.9, .24, C.GREEN, name="text_to_artifact")
    elif stage_index == 1:
        b.line(5.74, 3.7, 10.68, 3.7, C.GREEN, 1.8, name="structured_output_line")
        b.arrow(10.48, 3.57, .22, .26, C.GREEN, name="structured_output_arrow")
    elif stage_index < 6:
        b.line(7.98, 3.7, 10.68, 3.7, C.GREEN, 1.8, name="tool_to_artifact_line")
        b.arrow(10.48, 3.57, .22, .26, C.GREEN, name="tool_to_artifact_arrow")
    else:
        # S6 之后 PlanState 占据同一列，连线在节点两侧分段，避免压住文字。
        b.line(7.98, 3.7, 8.4, 3.7, C.GREEN, 1.8, name="tool_to_plan_line")
        b.line(10.11, 3.7, 10.68, 3.7, C.GREEN, 1.8, name="plan_to_artifact_line")
        b.arrow(10.48, 3.57, .22, .26, C.GREEN, name="plan_to_artifact_arrow")

    if stage_index >= 5:
        b.line(9.25, 3.02, 5.0, 3.02, C.GREEN, 1.2, True, name="state_context_line")
        b.text("latest State / Plan → Model Context", 5.52, 2.82, 2.55, .18,
               8.5, C.GREEN, True, PP_ALIGN.CENTER, name="state_context_label",
               allow_overlap=True)
    if stage_index >= 7:
        b.line(11.65, 3.08, 11.65, 4.6, C.RED, 1.8, name="artifact_review_line")
        b.line(11.65, 4.6, 10.1, 4.6, C.RED, 1.8, name="artifact_review_back")
        b.arrow(9.96, 4.47, .2, .26, C.RED, "left", "artifact_review_arrow")

    b.box("Recorder · events.jsonl · model_calls.jsonl · artifacts index", 3.65, 5.38, 5.95, .46,
          C.PALE, C.LIGHT, 11.5, C.GRAY, name="recorder")
    source_items = list(spec["sources"])
    source_body = "   ·   ".join(source_items)
    if len(source_items) >= 4:
        # 长源码映射拆成两行，投影环境中仍能看清模块名。
        source_body = "   ·   ".join(source_items[:2]) + "\n" + "   ·   ".join(source_items[2:])
    source_text = "本阶段新增源码：" + source_body
    b.rect(.72, 6.02, 11.86, .7, C.NAVY, C.NAVY, True, name="source_strip")
    b.text(source_text, .92, 6.08, 11.45, .54, 11.0, C.WHITE, False,
           PP_ALIGN.CENTER, font=FONT_MONO, name="source_text", allow_overlap=True)


def v2_bridge_visual(b: DeckBuilder, stage_index: int) -> None:
    """为八个输出页绘制不同但结构稳定的原生产物示意。"""

    b.rect(.72, 1.42, 6.0, 4.98, C.PALE, C.LIGHT, True, name="artifact_canvas",
           allow_overlap=True)
    if stage_index == 0:
        b.chip("draft.md", 1.02, 1.72, 1.05, C.GREEN, C.WHITE, name="draft_chip")
        b.text("市场概览\n\n第一部分……\n\n可能是第 3 页？\n\n数据来源：未明确", 1.08, 2.22, 4.05, 2.92,
               18, C.INK, valign=MSO_ANCHOR.TOP, name="draft_sample")
        for i, label in enumerate(["页数？", "字段？", "来源？"]):
            b.chip(label, 5.15, 2.35+i*.78, .92, C.PALE_RED, C.RED, name=f"draft_q_{i}")
    elif stage_index == 1:
        add_code_card(b, '{\n  "slides": [{\n    "id": "s1",\n    "title": "...",\n    "bullets": ["..."],\n    "source_ids": []\n  }]\n}', 1.0, 1.7, 4.28, 4.2, "DECKSPEC / JSON", "json_demo")
        b.box("deck.json", 5.45, 2.22, .98, .62, C.PALE_GREEN, C.GREEN, 13, C.GREEN, name="json_file")
        b.line(5.42, 3.05, 5.42, 4.68, C.RED, 2, True, name="json_break")
        b.text("data ≠ action", 5.22, 4.84, 1.2, .3, 11, C.RED, True, PP_ALIGN.CENTER, name="data_action")
    elif stage_index == 2:
        events = [("ToolCall", "search_web(query=…)", C.BLUE),
                  ("Runtime", "validate → execute", C.GRAY),
                  ("ToolResult", "call_id=42 · success", C.GREEN)]
        for i, (kind, content, color) in enumerate(events):
            y = 1.78+i*1.02
            b.chip(kind, 1.02, y+.18, 1.08, color, C.WHITE, name=f"event_kind_{i}")
            b.box(content, 2.32, y, 3.7, .7, C.WHITE, color, 14, C.INK, name=f"event_{i}")
        b.box("deck_v1.pptx", 2.2, 5.12, 3.05, .64, C.PALE_GREEN, C.GREEN, 17, C.GREEN, name="s2_ppt")
    elif stage_index == 3:
        steps = ["search", "read\npage", "Evidence", "compare", "PPT"]
        for i, label in enumerate(steps):
            x = .98+i*1.1
            color = C.ORANGE if i < 2 else C.GREEN
            b.box(label, x, 2.0, .86, .62, C.WHITE, color, 12, C.INK, name=f"fixed_step_{i}")
            if i < 4: b.arrow(x+.88, 2.2, .18, .2, C.MID, name=f"fixed_arrow_{i}")
        for i, label in enumerate(["sources.json", "evidence.json", "comparability.json", "full.pptx"]):
            b.box(label, 1.0+(i%2)*2.7, 3.4+(i//2)*.92, 2.35, .62,
                  C.PALE_GREEN, C.GREEN, 14, C.GREEN, name=f"s3_artifact_{i}")
    elif stage_index == 4:
        trace = [("ACTION", "搜索：AI 助手市场份额", C.BLUE),
                 ("OBS", "发现 MAU / 下载量混杂", C.GREEN),
                 ("ACTION", "改搜：中国大陆移动 App MAU", C.BLUE),
                 ("OBS", "历史数据仍有缺口", C.GREEN),
                 ("ACTION", "缩小范围并标注缺口", C.BLUE)]
        for i, (kind, content, color) in enumerate(trace):
            y = 1.66+i*.84
            b.chip(kind, .98, y+.13, .78, color, C.WHITE, name=f"trace_kind_{i}")
            b.box(content, 1.95, y, 4.38, .56, C.WHITE, color, 13.5, C.INK, name=f"trace_{i}")
    elif stage_index == 5:
        b.chip("state.json", 1.02, 1.72, 1.22, C.GREEN, C.WHITE, name="state_chip")
        fields = ["metric_scope", "source_ids", "open_questions", "completed_actions", "current_deck", "termination_reason"]
        for i, field in enumerate(fields):
            col, row = i%2, i//2
            b.box(field, 1.05+col*2.55, 2.42+row*.78, 2.25, .56,
                  C.WHITE, C.GREEN, 13.5, C.INK, name=f"state_field_{i}")
        b.box("Runtime updates only after real ToolResult", 1.42, 5.12, 4.45, .58,
              C.PALE_GREEN, C.PALE_GREEN, 14, C.GREEN, name="state_rule")
    elif stage_index == 6:
        for col, (title, color) in enumerate([("REVISION 1", C.BLUE), ("REVISION 2", C.GREEN)]):
            x = .98+col*2.85
            b.text(title, x, 1.72, 2.45, .28, 12, color, True, PP_ALIGN.CENTER, name=f"plan_title_{col}")
            plans = ["单一来源", "覆盖 12 个月", "直接画折线"] if col == 0 else ["寻找同口径", "缩小范围", "拆图 + 缺口"]
            for i, item in enumerate(plans):
                b.box(item, x, 2.22+i*.76, 2.45, .54, C.WHITE,
                      C.ORANGE if col and i else color, 13.5, C.INK, name=f"plan_{col}_{i}")
        b.arrow(3.53, 3.0, .28, .32, C.RED, name="replan_arrow")
        b.text("Observation", 3.18, 3.42, 1.0, .24, 10, C.RED, True, PP_ALIGN.CENTER, name="replan_obs")
    else:
        b.text("BEFORE", 1.02, 1.72, 2.0, .25, 11, C.RED, True, PP_ALIGN.CENTER, name="before_label")
        b.text("AFTER", 4.38, 1.72, 2.0, .25, 11, C.GREEN, True, PP_ALIGN.CENTER, name="after_label")
        b.rect(.98, 2.08, 2.28, 2.55, C.WHITE, C.RED, True, name="before_slide")
        b.text("市场份额\n32%\n来源缺失", 1.25, 2.64, 1.74, 1.25, 17, C.RED, True,
               PP_ALIGN.CENTER, name="before_content", allow_overlap=True)
        b.rect(4.25, 2.08, 2.28, 2.55, C.WHITE, C.GREEN, True, name="after_slide")
        b.text("样本内 MAU 份额\n32%\nsource_id · period", 4.52, 2.55, 1.74, 1.45,
               15, C.GREEN, True, PP_ALIGN.CENTER, name="after_content", allow_overlap=True)
        b.arrow(3.38, 2.82, .72, .42, C.GREEN, name="patch_arrow")
        b.chip("ReviewIssue → Patch", 2.56, 4.92, 2.42, C.PALE_RED, C.RED, name="patch_chip")


def v2_bridge_page(b: DeckBuilder, spec: dict, stage_index: int) -> None:
    """绘制统一的“当前输出—仍缺什么—下一步”桥接页面。"""

    b.new_slide(meta(f"s{stage_index}_output_gap", spec["bridge_title"], spec["stage"],
                     "3 分钟" if stage_index == 4 else "2 分钟", spec["message"],
                     spec["notes"], spec["next"]))
    v2_bridge_visual(b, stage_index)
    b.rect(7.02, 1.42, 5.56, 1.28, C.PALE_GREEN, C.GREEN, True, name="output_panel")
    b.chip("当前输出", 7.28, 1.66, 1.08, C.GREEN, C.WHITE, name="output_chip")
    b.text("  ·  ".join(spec["outputs"]), 7.34, 2.02, 4.94, .5, 13.8, C.INK,
           name="output_text", allow_overlap=True)
    b.rect(7.02, 2.98, 5.56, 2.15, C.PALE_RED, C.RED, True, name="gap_panel")
    b.chip("仍缺什么", 7.28, 3.22, 1.08, C.RED, C.WHITE, name="gap_chip")
    add_bullets(b, spec["gaps"], 7.3, 3.72, 4.94, .45, 13.8, C.INK, name="gap_list")
    b.rect(7.02, 5.34, 5.56, 1.06, C.NAVY, C.NAVY, True, name="next_panel")
    next_label = "NEXT" if stage_index < 7 else "BOUNDARY"
    b.chip(next_label, 7.28, 5.56, .92, C.WHITE, C.NAVY, name="next_chip")
    b.text(V2_NEXT_LABELS[stage_index], 8.32, 5.49, 4.02, .58, 13.8, C.WHITE, True,
           name="next_text", allow_overlap=True)


def v2_slide_01(b: DeckBuilder) -> None:
    """新版封面：先呈现骨架，不提前展开全部模块。"""

    b.new_slide(meta("cover", "从一次 API 调用到完整 Agent", "OPEN", "1 分钟",
                     "同一套系统连续增加八种能力。",
                     ["持续观察模型输出、下一步控制者和可检查产物。", "最终 PPT 是载体，主题是控制循环如何形成。"],
                     "先固定整个 Workshop 使用的同一个任务。"), show_progress=False)
    b.text("PHD WORKSHOP · AGENT FROM SCRATCH", .75, .58, 5.5, .3, 12, C.BLUE, True, name="cover_eyebrow")
    b.text("从一次 API 调用到\n完整 Agent", .75, 1.2, 7.2, 1.46, 37, C.NAVY, True,
           valign=MSO_ANCHOR.TOP, name="cover_title")
    b.text("同一个任务，八次增量：S0 → S7", .78, 2.84, 5.8, .42, 20, C.GRAY, name="cover_subtitle")
    columns = ["TASK", "CONTROLLER", "MODEL", "RUNTIME", "STATE / PLAN", "ARTIFACTS"]
    for i, label in enumerate(columns):
        x = .78+i*2.04
        fill = C.PALE_GREEN if i == 5 else C.PALE_BLUE
        line = C.GREEN if i == 5 else C.BLUE
        b.box(label, x, 4.02, 1.58, .7, fill, line, 12.5, C.NAVY, name=f"cover_col_{i}")
        if i < 5: b.arrow(x+1.65, 4.25, .24, .24, C.MID, name=f"cover_arrow_{i}")
    for i in range(8):
        b.chip(f"S{i}", 2.55+i*1.02, 5.38, .68,
               C.NAVY if i == 0 else C.PALE_BLUE, C.WHITE if i == 0 else C.NAVY,
               name=f"cover_stage_{i}")
    b.text("加了什么？   能输出什么？   还缺什么？", 2.6, 6.04, 8.15, .38,
           18, C.NAVY, True, PP_ALIGN.CENTER, name="cover_questions")
    b.text("API · Structured Output · Tools · Workflow · Agent Loop · State · Planning · Reflection",
           .8, 6.68, 11.7, .25, 11.5, C.GRAY, False, PP_ALIGN.CENTER, name="cover_keywords")


def v2_slide_02(b: DeckBuilder) -> None:
    """统一任务：固定输入和五页交付要求。"""

    b.new_slide(meta("task", "统一任务：完成一份真实市场调研 PPT", "TASK", "2 分钟",
                     "任务难点是研究、核对和修订，而不是生成文字。",
                     ["来源会混用 MAU、DAU、下载量和访问量。", "每个阶段处理同一个任务，便于比较新增能力。"],
                     "需要一组可观察标准判断系统处于哪个阶段。"), show_progress=False)
    b.rect(.75, 1.42, 4.42, 4.8, C.PALE_BLUE, C.LIGHT, True, name="task_card")
    b.chip("TASK", 1.02, 1.7, .76, C.NAVY, C.WHITE, name="task_chip")
    b.text("截至指定日期，制作 5 页中文 PPT，分析中国大陆消费端通用 AI 助手的竞争格局及最近 12 个月变化。",
           1.03, 2.18, 3.85, 1.22, 20, C.NAVY, True, valign=MSO_ANCHOR.TOP,
           name="task_statement", allow_overlap=True)
    add_bullets(b, ["5–6 款独立移动 App", "优先使用 MAU", "至少 3 个可比时间点", "关键数字可追溯"],
                1.03, 3.6, 3.7, .5, 16, name="task_rules")
    b.box("MAU ≠ DAU ≠ 下载量 ≠ 访问量", 1.02, 5.56, 3.88, .48,
          C.PALE_RED, C.PALE_RED, 15, C.RED, name="metric_warning", allow_overlap=True)
    b.arrow(5.46, 3.14, .72, .42, C.BLUE, name="task_to_output")
    b.text("Research\nVerify\nBuild", 5.35, 3.72, .92, 1.0, 13, C.BLUE, True,
           PP_ALIGN.CENTER, name="task_process")
    outputs = ["范围与结论", "最新竞争格局", "12 个月变化", "产品优劣势", "趋势与局限"]
    b.text("5-PAGE OUTPUT", 6.58, 1.46, 2.2, .28, 12, C.BLUE, True, name="output_title")
    for i, label in enumerate(outputs):
        b.box(f"{i+1:02d}   {label}", 6.58, 1.92+i*.82, 5.58, .62, C.WHITE, C.LIGHT,
              16.5, C.INK, name=f"task_output_{i}")


def v2_slide_03(b: DeckBuilder) -> None:
    """Stage 总览：阶段按新增控制机制划分。"""

    b.new_slide(meta("stage_contract", "八个 Stage 按“新增控制机制”划分", "MAP", "2 分钟",
                     "后一阶段复用前一阶段，只增加一个主要机制。",
                     ["S1 只输出 JSON，S2 才创建 PPT。", "S3 与 S4 的差异是控制权，不是工具集合。"],
                     "下一页说明八张结构图应该怎样阅读。"), show_progress=False)
    stages = [("S0", "API", "draft.md"), ("S1", "Structured Output", "deck.json"),
              ("S2", "Tools", "tool events"), ("S3", "Workflow", "full PPT"),
              ("S4", "Agent Loop", "agent trace"), ("S5", "State", "state.json"),
              ("S6", "Planning", "plan.json"), ("S7", "Reflection", "deck_v2 + review")]
    for i, (stage, mechanism, artifact) in enumerate(stages):
        row, col = divmod(i, 4); x, y = .82+col*3.08, 1.48+row*2.12
        color = C.ORANGE if stage == "S2" else (C.GREEN if stage in {"S5", "S6", "S7"} else C.BLUE)
        b.rect(x, y, 2.62, 1.5, C.PALE, color, True, name=f"map_{i}")
        b.chip(stage, x+.18, y+.18, .62, color, C.WHITE, name=f"map_stage_{i}")
        b.text(mechanism, x+.18, y+.63, 2.24, .38, 17, C.NAVY, True,
               name=f"map_mechanism_{i}", allow_overlap=True)
        b.text(artifact, x+.18, y+1.12, 2.24, .22, 11.5, C.GRAY,
               name=f"map_artifact_{i}", allow_overlap=True)
        if col < 3: b.arrow(x+2.7, y+.62, .22, .24, C.MID, name=f"map_arrow_{i}")
    b.box("后一阶段复用前一阶段，只增加一个主要机制", 3.22, 6.08, 6.9, .52,
          C.NAVY, C.NAVY, 17, C.WHITE, name="map_message")


def v2_slide_04(b: DeckBuilder) -> None:
    """固定读图规则：六列位置与增量高亮。"""

    b.new_slide(meta("diagram_legend", "同一张结构图，逐个 Stage 增量加工", "MAP", "2 分钟",
                     "位置不变，只观察新增高亮与新增数据流。",
                     ["浅色是继承能力，+ NEW 是本阶段新增。", "每个结构框都对应真实源码职责。"],
                     "从 S0 开始，此时系统只有一条直线。"), show_progress=False)
    labels = ["TASK", "CONTROLLER", "MODEL / CONTRACT", "RUNTIME / TOOLS", "STATE / PLAN / REVIEW", "ARTIFACTS"]
    for i, label in enumerate(labels):
        x = .72+i*2.03
        b.text(label, x, 1.48, 1.62, .24, 9.5, C.GRAY, True, PP_ALIGN.CENTER, name=f"legend_heading_{i}")
        kind = "tool" if i == 3 else ("data" if i in {4, 5} else "control")
        v2_new_node(b, ["config", "stage", "Model", "Tool", "State", "file.json"][i],
                    x, 2.0, 1.62, .76, kind, 0 if i != 3 else 1, 1, f"legend_node_{i}")
        if i < 5: b.arrow(x+1.7, 2.25, .2, .22, C.MID, name=f"legend_arrow_{i}")
    b.box("Recorder：从 S0 起贯穿所有 Stage，但不参与业务判断", 3.35, 3.22, 6.65, .5,
          C.PALE, C.LIGHT, 14, C.GRAY, name="legend_recorder")
    b.text("S0：直线", 1.02, 4.38, 2.1, .28, 13, C.BLUE, True, name="legend_s0")
    b.line(1.08, 5.02, 4.32, 5.02, C.BLUE, 2, name="legend_straight")
    b.arrow(4.1, 4.88, .24, .28, C.BLUE, name="legend_straight_arrow")
    b.text("S7：反馈回环", 7.0, 4.38, 2.3, .28, 13, C.GREEN, True, name="legend_s7")
    b.line(7.08, 5.02, 10.85, 5.02, C.BLUE, 2, name="legend_forward")
    b.line(10.85, 5.02, 10.85, 5.72, C.GREEN, 2, name="legend_loop_down")
    b.line(10.85, 5.72, 7.12, 5.72, C.GREEN, 2, name="legend_loop_back")
    b.arrow(6.92, 5.58, .22, .28, C.GREEN, "left", "legend_loop_arrow")
    b.chip("+ NEW", 8.38, 4.82, .72, C.NAVY, C.WHITE, name="legend_new")
    b.text("回环 = Observation 改变后续动作", 7.15, 6.02, 3.75, .3, 14, C.GREEN, True,
           PP_ALIGN.CENTER, name="legend_loop_label")


def v2_slide_21(b: DeckBuilder) -> None:
    """累计架构：标出所有模块首次出现的 Stage。"""

    b.new_slide(meta("cumulative_architecture", "完整系统：S0 的直线如何长成 S7 的闭环", "ALL", "1 分钟",
                     "完整 Agent 是多个清晰职责围绕控制循环协作。",
                     ["模型不直接执行，Tool 不决定下一步，Recorder 不参与业务判断。", "三条闭环分别是 Observation、State/Plan 回注和 Review/Patch。"],
                     "最后用控制权与输出矩阵回看八个阶段。"), show_progress=False)
    columns = [("TASK", .72), ("CONTROL", 2.35), ("MODEL / CONTRACT", 4.18),
               ("RUNTIME / TOOLS", 6.18), ("STATE / PLAN / REVIEW", 8.45), ("ARTIFACTS", 10.78)]
    for label, x in columns:
        b.text(label, x, 1.38, 1.65, .22, 9.5, C.GRAY, True, PP_ALIGN.CENTER, name=f"all_heading_{label}")
    nodes = [
        ("S0", "config.yaml", .72, 1.88, 1.42, .7, "control"),
        ("S4", "Agent Loop", 2.35, 1.88, 1.5, .7, "control"),
        ("S0", "OpenAIModel", 4.18, 1.88, 1.52, .7, "model"),
        ("S1", "DeckSpec", 4.18, 2.88, 1.52, .7, "data"),
        ("S2", "Runtime", 6.18, 1.88, 1.52, .7, "runtime"),
        ("S2", "Research\nTools", 6.05, 2.88, .92, .75, "tool"),
        ("S2", "PPT\nTools", 7.02, 2.88, .92, .75, "tool"),
        ("S5", "TaskState", 8.45, 1.88, 1.55, .66, "data"),
        ("S6", "PlanState", 8.45, 2.82, 1.55, .66, "data"),
        ("S7", "ReviewIssue", 8.45, 3.76, 1.55, .66, "issue"),
        ("S7", "deck_v2\nreview_v2", 10.78, 1.88, 1.8, .82, "data"),
    ]
    for i, (stage, text_value, x, y, w, h, kind) in enumerate(nodes):
        v2_new_node(b, text_value, x, y, w, h, kind, -1, 99, f"all_node_{i}")
        color = C.RED if stage == "S7" and kind == "issue" else (C.GREEN if kind == "data" else C.NAVY)
        b.chip(stage, x-.02, y-.18, .52, color, C.WHITE, name=f"all_stage_{i}")
    b.arrow(2.14, 2.1, .18, .22, C.BLUE, name="all_a1")
    b.arrow(3.92, 2.1, .2, .22, C.BLUE, name="all_a2")
    b.arrow(5.76, 2.1, .34, .22, C.BLUE, name="all_a3")
    b.line(7.7, 4.6, 4.92, 4.6, C.GREEN, 2.2, name="all_observation")
    b.arrow(4.7, 4.47, .24, .26, C.GREEN, "left", "all_observation_arrow")
    b.text("ToolResult / Observation", 5.25, 4.3, 2.05, .2, 9.5, C.GREEN, True,
           PP_ALIGN.CENTER, name="all_observation_label")
    b.line(9.2, 4.52, 5.0, 5.05, C.GREEN, 1.4, True, name="all_context_return")
    b.text("State + Plan → Model Context", 5.7, 4.82, 2.75, .2, 9.5, C.GREEN, True,
           PP_ALIGN.CENTER, name="all_context_label")
    b.line(11.7, 2.82, 11.7, 4.2, C.RED, 1.8, name="all_review_down")
    b.line(11.7, 4.2, 10.12, 4.2, C.RED, 1.8, name="all_review_back")
    b.arrow(9.94, 4.07, .2, .26, C.RED, "left", "all_review_arrow")
    b.box("Recorder · events · model calls · artifact index", 3.68, 5.35, 5.95, .44,
          C.PALE, C.LIGHT, 11.5, C.GRAY, name="all_recorder")
    artifacts = ["draft", "JSON", "tools", "evidence", "agent", "state", "plan", "review"]
    for i, label in enumerate(artifacts):
        x = .82+i*1.55
        b.box(f"S{i}\n{label}", x, 6.0, 1.16, .62, C.PALE_BLUE if i < 7 else C.PALE_GREEN,
              C.BLUE if i < 7 else C.GREEN, 12.5, C.NAVY, name=f"all_timeline_{i}")
        if i < 7: b.arrow(x+1.2, 6.2, .2, .2, C.MID, name=f"all_timeline_arrow_{i}")


def v2_slide_22(b: DeckBuilder) -> None:
    """总结矩阵：控制权、可观察输出与关键缺口。"""

    b.new_slide(meta("takeaway_matrix", "判断 Agent，看控制权与反馈链", "END", "0.5 分钟",
                     "环境反馈是否改变下一步，是可观察的核心判断。",
                     ["确定性规则留给代码，开放式选择交给模型。", "实际系统不应追求最大自治。"],
                     "进入提问或打开真实运行目录。"), show_progress=False)
    data = [
        ["Stage", "谁决定下一步", "可观察输出", "关键缺口"],
        ["S0–S1", "宿主程序", "文本 / JSON", "无外部执行"],
        ["S2", "宿主控制回合", "Tool 轨迹 / 示例 PPT", "无端到端流程"],
        ["S3", "固定 Python", "完整 PPT", "路径不能动态改变"],
        ["S4", "模型依据 Observation", "Agent 轨迹", "无显式任务事实"],
        ["S5–S6", "模型 + Runtime 校验", "State / Plan / PPT", "无产物级反馈"],
        ["S7", "模型依据 ReviewIssue", "修订版本与复查", "Workshop 范围完成"],
    ]
    b.table(data, .82, 1.48, 11.7, 4.62, [1.1, 2.4, 2.6, 2.5], True, 13.5, "takeaway_table")
    b.box("判断 Agent，不看它做了多少事；看环境反馈是否改变了下一步",
          1.92, 6.28, 9.5, .52, C.NAVY, C.NAVY, 18, C.WHITE, name="takeaway_message")


def v2_appendix_a(b: DeckBuilder) -> None:
    """备用 A：结构模块与源码目录一一对应。"""

    b.new_slide(meta("appendix_source_map", "结构框中的模块如何对应源码目录", "APP", "备用",
                     "每个框都能落到明确的源码职责。",
                     ["Stage 负责控制流，Tool 只执行一次动作。", "Runtime 是唯一工具分派与循环实现。"]), show_progress=False)
    tree = "src/\n├─ run.py\n├─ stages/\n├─ core/model.py\n├─ core/contracts.py\n├─ core/runtime.py\n├─ core/context.py\n├─ core/recorder.py\n└─ tools/"
    add_code_card(b, tree, .82, 1.45, 4.55, 4.95, "SOURCE TREE", "source_tree")
    mappings = [("Stage Controller", "run.py + stages/", C.NAVY),
                ("Model / Contract", "model.py + contracts.py", C.BLUE),
                ("Runtime / Tools", "runtime.py + tools/", C.ORANGE),
                ("State / Plan", "context.py + contracts.py", C.GREEN),
                ("Recorder", "recorder.py", C.GRAY)]
    for i, (left, right, color) in enumerate(mappings):
        y = 1.56+i*.93
        b.box(left, 5.88, y, 2.25, .58, C.WHITE, color, 14, C.INK, name=f"source_map_left_{i}")
        b.arrow(8.24, y+.18, .42, .22, color, name=f"source_map_arrow_{i}")
        b.box(right, 8.8, y, 3.35, .58, C.PALE, color, 13.5, C.INK, name=f"source_map_right_{i}")
    b.box("结构图不是概念拼贴：每个节点都有单一源码职责", 5.92, 6.02, 6.2, .52,
          C.PALE_GREEN, C.PALE_GREEN, 16, C.GREEN, name="source_map_message")


def v2_appendix_b(b: DeckBuilder) -> None:
    """备用 B：核心数据契约。"""

    b.new_slide(meta("appendix_contracts", "核心数据契约怎样串起系统", "APP", "备用",
                     "Schema 约束格式，来源和检查约束可信度。",
                     ["结构化数据不自动等于真实数据。", "Evidence 可比性必须检查 metric、period、geography 和 platform。"]), show_progress=False)
    pairs = [("ToolCall", "ToolResult", "Runtime"), ("Source", "Evidence", "Model + Contract"),
             ("DeckSpec", "PPT", "Presentation Tool"), ("TaskState + Plan", "Model Context", "Runtime"),
             ("ReviewIssue", "DeckPatch", "Agent + Runtime")]
    for i, (left, right, owner) in enumerate(pairs):
        y = 1.42+i*.96
        b.box(left, .92, y, 3.12, .62, C.PALE_BLUE if i != 4 else C.PALE_RED,
              C.BLUE if i != 4 else C.RED, 15, C.NAVY, name=f"contract_left_{i}")
        b.arrow(4.28, y+.2, .62, .24, C.GREEN, name=f"contract_arrow_{i}")
        b.chip(owner, 5.0, y+.14, 1.72, C.PALE, C.GRAY, name=f"contract_owner_{i}")
        b.arrow(6.86, y+.2, .62, .24, C.GREEN, name=f"contract_arrow_b_{i}")
        b.box(right, 7.72, y, 3.68, .62, C.PALE_GREEN if i != 4 else C.PALE_ORANGE,
              C.GREEN if i != 4 else C.ORANGE, 15, C.NAVY, name=f"contract_right_{i}")
    b.box("Schema 保证接口边界；Evidence 与 Review 保证可追溯性", 2.75, 6.3, 7.85, .5,
          C.PALE_GREEN, C.PALE_GREEN, 16, C.GREEN, name="contract_message")


def v2_appendix_c(b: DeckBuilder) -> None:
    """备用 C：运行目录支持复现和录屏。"""

    b.new_slide(meta("appendix_artifacts", "一次运行目录如何支持复现与课堂演示", "APP", "备用",
                     "可回放文件替代隐藏推理。",
                     ["每次运行创建新目录，禁止静默覆盖。", "课堂展示请求、工具事件、状态和版本文件。"]), show_progress=False)
    tree = "runs/<run_id>/\n├─ config.snapshot.yaml\n├─ events.jsonl\n├─ model_calls.jsonl\n├─ sources.json / evidence.json\n├─ state.json / plan.json\n├─ deck_v1.pptx / deck_v2.pptx\n└─ review_v1.json / review_v2.json"
    add_code_card(b, tree, .82, 1.45, 5.6, 4.95, "RUN ARTIFACTS", "run_tree")
    uses = [("结构证据", "模块加入了什么", C.BLUE), ("轨迹证据", "Observation 改变 Action", C.GREEN),
            ("版本证据", "Before / Patch / After", C.RED)]
    for i, (title, body, color) in enumerate(uses):
        y = 1.68+i*1.35
        b.chip(title, 7.0, y+.22, 1.12, color, C.WHITE, name=f"artifact_use_{i}")
        b.box(body, 8.36, y, 3.68, .82, C.WHITE, color, 16, C.INK, name=f"artifact_body_{i}")
    b.box("不展示隐藏思维链，只展示可观察输入、动作、结果与状态", 6.95, 5.98, 5.18, .56,
          C.PALE_RED, C.PALE_RED, 14.5, C.RED, name="artifact_rule")


def v2_appendix_d(b: DeckBuilder) -> None:
    """备用 D：Agent Loop 的 Runtime 边界。"""

    b.new_slide(meta("appendix_limits", "Agent Loop 必须由 Runtime 设置边界", "APP", "备用",
                     "模型不能自行宣布成功，也不能无限循环。",
                     ["Runtime 检查 PPT 是否实际存在。", "达到任一预算上限都产生明确结束原因。"]), show_progress=False)
    b.box("MODEL\nchooses Action", 4.25, 2.12, 2.2, .88, C.PALE_BLUE, C.BLUE, 17, C.NAVY, name="limit_model")
    b.hex("Runtime\nexecutes Tool", 7.45, 2.12, 2.25, .88, C.PALE_ORANGE, C.ORANGE, 15, "limit_runtime")
    b.box("Observation", 5.85, 4.0, 2.2, .72, C.PALE_GREEN, C.GREEN, 17, C.GREEN, name="limit_obs")
    b.arrow(6.62, 2.38, .62, .28, C.BLUE, name="limit_a1")
    b.line(8.55, 3.14, 8.55, 4.38, C.GREEN, 2, name="limit_a2")
    b.line(8.55, 4.38, 8.15, 4.38, C.GREEN, 2, name="limit_a3")
    b.line(5.7, 4.38, 5.05, 4.38, C.GREEN, 2, name="limit_a4")
    b.arrow(4.82, 3.18, .3, .52, C.GREEN, "up", "limit_a5")
    limits = ["max_model_calls", "max_tool_calls", "max_agent_steps", "max_elapsed_seconds", "max_review_rounds"]
    for i, label in enumerate(limits):
        b.chip(label, .78, 1.48+i*.78, 2.75, C.PALE_RED, C.RED, name=f"runtime_limit_{i}")
    outcomes = ["completed", "incomplete", "limit_reached", "error"]
    for i, label in enumerate(outcomes):
        b.chip(label, 10.25, 1.7+i*.9, 2.05,
               C.PALE_GREEN if i == 0 else C.PALE_RED, C.GREEN if i == 0 else C.RED,
               name=f"runtime_outcome_{i}")
    b.box("Runtime 验证真实交付物，并决定何时停止", 3.75, 5.8, 6.25, .58,
          C.NAVY, C.NAVY, 17, C.WHITE, name="runtime_boundary_message")


def build_v2_deck(b: DeckBuilder) -> None:
    """按新版设计一次生成 22 页主讲与 4 页备用页。"""

    v2_slide_01(b); v2_slide_02(b); v2_slide_03(b); v2_slide_04(b)
    for stage_index, spec in enumerate(V2_STAGES):
        v2_architecture_page(b, spec, stage_index)
        v2_bridge_page(b, spec, stage_index)
    v2_slide_21(b); v2_slide_22(b)
    v2_appendix_a(b); v2_appendix_b(b); v2_appendix_c(b); v2_appendix_d(b)


SLIDE_BUILDERS = [build_v2_deck]


def main() -> int:
    """生成 PPTX、讲者备注、预览图与质量检查报告。"""

    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    builder = DeckBuilder()
    for build_slide in SLIDE_BUILDERS:
        build_slide(builder)
    report = builder.finish()
    print(f"已生成：{PPTX_PATH}")
    print(f"页面数：{report['slide_count']}；检查问题：{report['issue_count']}")
    if report["office_render"]["rendered"]:
        print(f"真实渲染目录：{RENDER_DIR}")
    else:
        print(f"真实渲染未执行：{report['office_render']['reason']}")
    print(f"预览目录：{PREVIEW_DIR}")
    print(f"质量报告：{REPORT_PATH}")
    return 0 if report["issue_count"] == 0 else 2


if __name__ == "__main__":
    sys.exit(main())
