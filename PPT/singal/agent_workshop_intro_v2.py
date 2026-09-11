#!/usr/bin/env python3
"""
Agent Workshop opening slides generator.

main() only calls build_intro_slides().
The design follows the existing deck:
- white background
- navy titles
- blue = model/decision
- orange = tools/actions
- green = feedback/state
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


OUTPUT = "agent_workshop_intro_v2.pptx"


class C:
    NAVY = "17324D"
    BLUE = "315B7D"
    ORANGE = "D97706"
    GREEN = "2F855A"
    GRAY = "64748B"
    LIGHT = "E2E8F0"
    PALE_BLUE = "EAF1F7"
    PALE_GREEN = "EAF6EF"
    PALE_ORANGE = "FFF4E6"
    PALE_RED = "FCEDEC"
    RED = "C2413B"


def rgb(c):
    return RGBColor.from_string(c)


def add_text(slide, t, x, y, w, h, size=18, color=C.NAVY, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    p = box.text_frame.paragraphs[0]
    p.text = t
    for r in p.runs:
        r.font.name = "Aptos"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
    return box


def add_card(slide, t, x, y, w, h, fill, line):
    s = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.color.rgb = rgb(line)
    add_text(slide, t, x+.1, y+.1, w-.2, h-.2, 16, C.NAVY, True)


def title(slide, t):
    add_text(slide, t, .7, .35, 12, .5, 30, C.NAVY, True)


def slide_products(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Are We Already Using Agents Every Day?")

    add_text(s,
             "ChatGPT, Claude Code, and AI assistants look similar, but they have different levels of autonomy.",
             1, 1.3, 11, .5, 18, C.GRAY)

    cards = [
        ("ChatGPT\n\nAsk → Answer", C.PALE_BLUE, C.BLUE),
        ("Claude Code\n\nPlan → Edit → Test", C.PALE_GREEN, C.GREEN),
        ("AI Assistant\n\nSearch → Recommend", C.PALE_ORANGE, C.ORANGE),
    ]

    for i, (txt, fill, line) in enumerate(cards):
        add_card(s, txt, 1+i*4, 2.5, 3, 1.8, fill, line)

    add_text(s, "Question: Which one is an Agent? Why?",
             2, 5.5, 9, .5, 24, C.NAVY, True)


def slide_llm_agent(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "LLM Generates. Agents Act.")

    add_card(s, "Prompt", 1, 3, 1.5, .8, C.PALE_BLUE, C.BLUE)
    add_card(s, "LLM", 3.2, 3, 1.5, .8, C.PALE_BLUE, C.BLUE)
    add_card(s, "Text Response", 5.4, 3, 2, .8, C.PALE_GREEN, C.GREEN)

    add_card(s, "Goal", 8.2, 3, 1.3, .8, C.PALE_BLUE, C.BLUE)
    add_card(s, "Decision", 10, 3, 1.5, .8, C.PALE_BLUE, C.BLUE)
    add_card(s, "Action", 11.8, 3, 1, .8, C.PALE_ORANGE, C.ORANGE)

    add_text(s,
             "LLM: generate information\n\nAgent: decide and execute actions",
             2, 5.1, 9, .8, 22, C.NAVY, True)


def slide_boundary(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "A Language Model Only Speaks Language")

    add_card(s, "Text Input", 1, 3, 2, 1, C.PALE_BLUE, C.BLUE)
    add_card(s, "Language Model", 5, 3, 2.5, 1, C.PALE_BLUE, C.BLUE)
    add_card(s, "Text Output", 9.5, 3, 2, 1, C.PALE_GREEN, C.GREEN)

    add_card(s,
             "Without external systems:\n"
             "✗ Search\n✗ Code execution\n✗ File modification\n✗ Database access",
             3.5, 4.8, 6, 1.4, C.PALE_RED, C.RED)

    add_text(s,
             "Agent connects language models with the external world.",
             2, 6.4, 9, .4, 20, C.NAVY, True)


def slide_evolution(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "We Don't Start with an Agent. We Build One.")

    stages = [
        "API", "Structured Output", "Tools", "Workflow",
        "Agent Loop", "State", "Planning", "Reflection"
    ]

    for i, stage in enumerate(stages):
        x = .8 + (i % 4) * 3.1
        y = 1.8 + (i // 4) * 1.8
        add_card(s, stage, x, y, 2.5, .9,
                 C.PALE_GREEN if i >= 4 else C.PALE_BLUE,
                 C.GREEN if i >= 4 else C.BLUE)

    add_text(s,
             "Today: building an Agent step by step, starting from the first API call.",
             1, 6.2, 11, .4, 20, C.NAVY, True)


def build_intro_slides():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_products(prs)
    slide_llm_agent(prs)
    slide_boundary(prs)
    slide_evolution(prs)

    prs.save(OUTPUT)


def main():
    build_intro_slides()


if __name__ == "__main__":
    main()
