"""实现 PPT 构建、渲染和结构化页面修订。首次使用：S2。"""

from __future__ import annotations

from pathlib import Path
import math
import shutil
import subprocess
from typing import Any

from core.config import AppConfig
from core.contracts import apply_deck_patches, validate_deck


class PresentationToolError(RuntimeError):
    """表示需要作为 Observation 返回的演示文稿工具错误。首次使用：S2。"""


# 全角字符从此码位开始，按 1 个宽度单位估算；半角按 0.55 估算。
FULL_WIDTH_THRESHOLD = 0x2E80
# 单行实际高度约为字号的倍数，用于把行数换算成需要的英寸高度。
LINE_SPACING = 1.25
# 相邻 bullet 之间的段后间距（英寸），与 _add_bullets 里设置的 space_after 保持一致。
BULLET_GAP_INCHES = 4.0 / 72.0


def _text_width_units(text: str) -> float:
    """估算文本的显示宽度单位，用于预测自动换行后的行数。首次使用：S2。"""

    return sum(1.0 if ord(char) >= FULL_WIDTH_THRESHOLD else 0.55 for char in text)


def _estimate_line_count(text: str, width_inches: float, font_size: int) -> int:
    """按框宽估算自动换行后的行数；每个换行符单独起一行。首次使用：S2。"""

    capacity = max(4.0, (width_inches * 72.0) / float(font_size))
    return sum(
        max(1, math.ceil(_text_width_units(segment) / capacity)) for segment in text.split("\n")
    )


def _fit_font_size(
    text: str,
    width_inches: float,
    height_inches: float,
    preferred_size: int,
    minimum_size: int,
) -> int:
    """从首选字号逐磅下调，直到文本能放进文本框为止。首次使用：S2。"""

    for size in range(preferred_size, minimum_size - 1, -1):
        needed = _estimate_line_count(text, width_inches, size) * size * LINE_SPACING / 72.0
        if needed <= height_inches:
            return size
    return minimum_size


def _theme_settings(config: AppConfig) -> dict[str, Any]:
    """合并当前主题与预设默认值，供各版式渲染时读取。首次使用：S3。"""

    presentation = config.section("presentation")
    theme_name = str(presentation.get("theme", "")).strip()
    themes = presentation.get("themes") or {}
    theme = themes.get(theme_name, {}) if isinstance(themes, dict) else {}
    if not isinstance(theme, dict):
        theme = {}

    def pick(key: str, fallback: Any) -> Any:
        """优先取主题值，其次取 presentation 顶层值，最后用内置默认。"""

        if theme.get(key) is not None:
            return theme[key]
        if presentation.get(key) is not None:
            return presentation[key]
        return fallback

    return {
        "font_family": str(pick("font_family", "Noto Sans CJK SC")),
        "title_size": int(pick("title_size", 26)),
        "body_size": int(pick("body_size", 18)),
        "chart_body_size": int(pick("chart_body_size", 14)),
        "table_size": int(pick("table_size", 12)),
        "small_size": int(pick("small_size", 9)),
        "title_color": str(pick("title_color", "1F1F1F")),
        "text_color": str(pick("text_color", "333333")),
        "muted_color": str(pick("muted_color", "7F7F7F")),
        "accent": str(pick("accent", "1F4E79")),
        "accent_text": str(pick("accent_text", "FFFFFF")),
    }


def _style_paragraph(
    paragraph: Any,
    font_name: str,
    font_size: int,
    bold: bool = False,
    color: str | None = None,
) -> None:
    """统一设置字号、粗细与中英文 typeface，并给中文补上东亚字体节点。首次使用：S2。"""

    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    font = paragraph.font
    font.name = font_name
    font.size = Pt(font_size)
    font.bold = bold
    if color is not None:
        font.color.rgb = RGBColor.from_string(color)
    # python-pptx 的 font.name 只写 <a:latin>；中文实际由 <a:ea> 决定。
    # 这里在 defRPr 里把 latin 之后的 ea/cs 一并补上，否则中文会落到主题默认字体。
    r_pr = paragraph._p.get_or_add_pPr().get_or_add_defRPr()
    latin = r_pr.find(qn("a:latin"))
    if latin is None:
        return
    previous = latin
    for tag in ("a:ea", "a:cs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = r_pr.makeelement(qn(tag), {})
            previous.addnext(node)
        node.set("typeface", font_name)
        previous = node


def _add_textbox(
    slide: Any,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    font_name: str,
    font_size: int,
    bold: bool = False,
    color: str | None = None,
    anchor: str = "t",
) -> None:
    """按照英寸坐标添加样式统一的文本框。首次使用：S2。"""

    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.util import Inches

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    # python-pptx 的 add_textbox() 默认写入 <a:bodyPr wrap="none"><a:spAutoFit/>：
    # 文本不换行，会向右溢出文本框，只有手动微调形状时才重新排版。
    # 这里显式开启自动换行、固定形状尺寸，并把内边距归零，使坐标即文本起点。
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    _style_paragraph(paragraph, font_name, font_size, bold, color)


def _add_bullets(
    slide: Any,
    bullets: list[str],
    left: float,
    top: float,
    width: float,
    height: float,
    theme: dict[str, Any],
    font_size: int,
) -> None:
    """把每一条要点画成独立段落，显式控制行距和段距。首次使用：S2。"""

    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.clear()
    for index, item in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"• {item}"
        paragraph.line_spacing = LINE_SPACING
        if index < len(bullets) - 1:
            paragraph.space_after = Pt(4)
        _style_paragraph(paragraph, theme["font_family"], font_size, False, theme["text_color"])


CHART_TYPES = {"bar", "line"}


def _chart_is_renderable(chart: Any) -> bool:
    """判断页面图表是否可以真正渲染出图形。首次使用：S3。"""

    if not isinstance(chart, dict):
        return False
    if str(chart.get("type", "none")) not in CHART_TYPES:
        return False
    return bool(chart.get("categories")) and bool(chart.get("series"))


def _add_chart(
    slide: Any,
    chart_data: dict[str, Any],
    left: float,
    top: float,
    width: float,
    height: float,
    font_name: str,
) -> None:
    """用 python-pptx 原生图表渲染页面数据，保持图表在 PPT 内可编辑。首次使用：S3。"""

    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt

    chart_type = (
        XL_CHART_TYPE.COLUMN_CLUSTERED if chart_data["type"] == "bar" else XL_CHART_TYPE.LINE_MARKERS
    )
    data = CategoryChartData()
    data.categories = list(chart_data["categories"])
    for series in chart_data["series"]:
        data.add_series(str(series["name"]), tuple(series["values"]))
    chart = slide.shapes.add_chart(
        chart_type, Inches(left), Inches(top), Inches(width), Inches(height), data
    ).chart

    title = str(chart_data.get("title", "")).strip()
    chart.has_title = bool(title)
    if title:
        chart.chart_title.text_frame.text = title
        for paragraph in chart.chart_title.text_frame.paragraphs:
            paragraph.font.size = Pt(12)
            paragraph.font.name = font_name

    chart.has_legend = len(chart_data["series"]) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(10)
        chart.legend.font.name = font_name
    chart.font.size = Pt(10)
    chart.font.name = font_name

    unit = str(chart_data.get("unit", "")).strip()
    if unit == "%":
        chart.value_axis.tick_labels.number_format = '0.0"%"'
    if unit:
        try:
            chart.value_axis.has_title = True
            chart.value_axis.axis_title.text_frame.text = unit
            chart.value_axis.axis_title.text_frame.paragraphs[0].font.size = Pt(10)
            chart.value_axis.axis_title.text_frame.paragraphs[0].font.name = font_name
        except (AttributeError, ValueError):
            # 老版本 python-pptx 没有 axis_title，此时只保留坐标轴刻度。
            chart.value_axis.has_title = False


def _draw_title(slide: Any, slide_data: dict[str, Any], theme: dict[str, Any]) -> None:
    """绘制统一的页面标题，并在过长时自动缩小字号。首次使用：S3。"""

    size = _fit_font_size(slide_data["title"], 11.7, 0.75, int(theme["title_size"]), 16)
    _add_textbox(
        slide,
        slide_data["title"],
        0.7,
        0.45,
        11.7,
        0.75,
        theme["font_family"],
        size,
        True,
        theme["title_color"],
    )


def _draw_bullets(
    slide: Any,
    bullets: list[str],
    left: float,
    top: float,
    width: float,
    height: float,
    theme: dict[str, Any],
    preferred_size: int,
    stat: dict[str, Any],
) -> None:
    """绘制一段要点，按内容高度收紧文本框，并记录是否缩小了字号。首次使用：S3。"""

    if not bullets:
        return
    text = "\n".join(f"• {item}" for item in bullets)
    size = _fit_font_size(text, width, height, preferred_size, 10)
    if size != preferred_size:
        stat["shrunk_here"] = True
    line_count = sum(_estimate_line_count(f"• {item}", width, size) for item in bullets)
    needed_height = line_count * size * LINE_SPACING / 72.0 + (len(bullets) - 1) * BULLET_GAP_INCHES
    _add_bullets(slide, bullets, left, top, width, min(height, needed_height), theme, size)


def _render_bullets(slide: Any, slide_data: dict[str, Any], theme: dict[str, Any], stat: dict[str, Any]) -> None:
    """单栏要点版式；有图表时正文让出右半版面。首次使用：S3。"""

    if _chart_is_renderable(slide_data.get("chart")):
        _draw_bullets(
            slide,
            slide_data["bullets"],
            0.9,
            1.45,
            5.9,
            4.9,
            theme,
            int(theme["chart_body_size"]),
            stat,
        )
        _add_chart(slide, slide_data["chart"], 7.0, 1.5, 5.6, 4.8, theme["font_family"])
        stat["charts"] += 1
    else:
        _draw_bullets(
            slide,
            slide_data["bullets"],
            0.9,
            1.45,
            11.5,
            4.9,
            theme,
            int(theme["body_size"]),
            stat,
        )


def _render_two_column(
    slide: Any, slide_data: dict[str, Any], theme: dict[str, Any], stat: dict[str, Any]
) -> None:
    """左右两栏要点版式，用于并列对比两组信息。首次使用：S3。"""

    bullets = slide_data["bullets"]
    if len(bullets) < 2:
        _render_bullets(slide, slide_data, theme, stat)
        return
    middle = (len(bullets) + 1) // 2
    for index, chunk in enumerate((bullets[:middle], bullets[middle:])):
        _draw_bullets(
            slide,
            chunk,
            0.9 + index * 6.0,
            1.45,
            5.6,
            4.9,
            theme,
            int(theme["body_size"]),
            stat,
        )


def _render_comparison_table(
    slide: Any, slide_data: dict[str, Any], theme: dict[str, Any], stat: dict[str, Any]
) -> None:
    """对比表格版式；缺少表格数据时回退为要点版式，避免出现空白页。首次使用：S3。"""

    table = slide_data.get("table") or {}
    if not (table.get("columns") and table.get("rows")):
        _render_bullets(slide, slide_data, theme, stat)
        return

    caption = str(table.get("caption", "")).strip()
    table_top = 1.45
    if caption:
        _add_textbox(
            slide,
            caption,
            0.9,
            1.3,
            11.5,
            0.4,
            theme["font_family"],
            int(theme["small_size"]) + 2,
            color=theme["muted_color"],
        )
        table_top = 1.85
    rows = len((table.get("rows") or []))
    height = min(4.4, 0.36 * (rows + 1) + 0.2)
    if _add_table(slide, table, 0.9, table_top, 11.5, height, theme):
        stat["tables"] += 1


def _render_chart_focus(
    slide: Any, slide_data: dict[str, Any], theme: dict[str, Any], stat: dict[str, Any]
) -> None:
    """图表为主、要点为辅的版式；没有图表数据时回退为要点版式。首次使用：S3。"""

    if not _chart_is_renderable(slide_data.get("chart")):
        _render_bullets(slide, slide_data, theme, stat)
        return
    _add_chart(slide, slide_data["chart"], 0.9, 1.4, 11.5, 3.6, theme["font_family"])
    stat["charts"] += 1
    _draw_bullets(
        slide,
        slide_data["bullets"][:3],
        0.9,
        5.2,
        11.5,
        1.5,
        theme,
        int(theme["chart_body_size"]),
        stat,
    )


def _add_table(
    slide: Any,
    table: dict[str, Any],
    left: float,
    top: float,
    width: float,
    height: float,
    theme: dict[str, Any],
) -> bool:
    """在指定区域内渲染原生表格，失败时返回 False。首次使用：S3。"""

    columns = [str(item) for item in (table.get("columns") or [])]
    rows = [[str(cell) for cell in row] for row in (table.get("rows") or [])]
    if not columns or not rows:
        return False

    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    shape = slide.shapes.add_table(
        len(rows) + 1, len(columns), Inches(left), Inches(top), Inches(width), Inches(height)
    )
    grid = shape.table
    for column_index, title in enumerate(columns):
        grid.cell(0, column_index).text = title
    for row_index, row in enumerate(rows, start=1):
        for column_index, cell_text in enumerate(row):
            grid.cell(row_index, column_index).text = cell_text

    for row_index in range(len(rows) + 1):
        for column_index in range(len(columns)):
            cell = grid.cell(row_index, column_index)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            if row_index == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(str(theme["accent"]))
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(int(theme["table_size"]))
                paragraph.font.name = theme["font_family"]
                paragraph.font.bold = row_index == 0
                if row_index == 0:
                    paragraph.font.color.rgb = RGBColor.from_string(str(theme["accent_text"]))
    return True


def _add_callout(
    slide: Any,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    theme: dict[str, Any],
) -> None:
    """渲染带左侧强调色条的一句话结论。首次使用：S3。"""

    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.08), Inches(height))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(str(theme["accent"]))
    bar.line.fill.background()
    _add_textbox(
        slide,
        text,
        left + 0.22,
        top,
        width - 0.22,
        height,
        theme["font_family"],
        int(theme["chart_body_size"]),
        True,
        theme["accent"],
        "m",
    )


def _add_kpi(
    slide: Any,
    value: str,
    label: str,
    left: float,
    top: float,
    width: float,
    height: float,
    theme: dict[str, Any],
    emphasis: str,
) -> None:
    """渲染大号强调数字 + 小标签，用于关键指标。首次使用：S3。"""

    value_size = {"high": 30, "medium": 26, "low": 20}[emphasis]
    _add_textbox(
        slide, value, left, top, width, height - 0.28, theme["font_family"], value_size, True, theme["accent"]
    )
    if label:
        _add_textbox(
            slide, label, left, top + height - 0.28, width, 0.28, theme["font_family"],
            int(theme["small_size"]) + 2, False, theme["muted_color"],
        )


def _element_height(element: dict[str, Any], width: float, theme: dict[str, Any]) -> float:
    """估算单个元素需要的高度，供自定义布局的流式排版使用。首次使用：S3。"""

    element_type = element["type"]
    if element_type == "callout":
        return 0.55
    if element_type == "kpi":
        return 0.95
    if element_type == "bullets":
        size = int(theme["body_size"])
        lines = sum(_estimate_line_count(f"• {item}", width, size) for item in element["items"])
        return lines * size * LINE_SPACING / 72.0 + (len(element["items"]) - 1) * BULLET_GAP_INCHES
    if element_type == "chart":
        return 2.7
    if element_type == "table":
        return min(2.8, 0.34 * (len(element["table"].get("rows") or []) + 1) + 0.3)
    return 0.6


def _render_element(
    slide: Any,
    element: dict[str, Any],
    left: float,
    top: float,
    width: float,
    height: float,
    theme: dict[str, Any],
    stat: dict[str, Any],
) -> None:
    """按元素类型分派渲染。首次使用：S3。"""

    element_type = element["type"]
    if element_type == "callout":
        _add_callout(slide, element["text"], left, top, width, height, theme)
    elif element_type == "kpi":
        _add_kpi(slide, element["value"], element["label"], left, top, width, height, theme, element["emphasis"])
    elif element_type == "bullets":
        _add_bullets(slide, element["items"], left, top, width, height, theme, int(theme["body_size"]))
    elif element_type == "chart":
        _add_chart(slide, element["chart"], left, top, width, height, theme["font_family"])
        stat["charts"] += 1
    elif element_type == "table":
        if _add_table(slide, element["table"], left, top, width, height, theme):
            stat["tables"] += 1


def _render_custom(slide: Any, slide_data: dict[str, Any], theme: dict[str, Any], stat: dict[str, Any]) -> None:
    """自定义布局：按元素列表做流式排版，模型决定“有什么”，代码决定几何。首次使用：S3。"""

    elements = slide_data.get("elements") or []
    if not elements:
        _render_bullets(slide, slide_data, theme, stat)
        return
    placements, overflow = simulate_custom_flow(elements, theme)
    for element, x, y, width, height in placements:
        _render_element(slide, element, x, y, width, height, theme, stat)
    if overflow:
        stat["custom_overflow"] = True


def _custom_body_area() -> tuple[float, float, float, float]:
    """返回自定义布局的正文区：left, top, width, bottom（bottom 留出来源行空间）。首次使用：S3。"""

    return 0.9, 1.45, 11.5, 6.55


def simulate_custom_flow(
    elements: list[dict[str, Any]],
    theme: dict[str, Any],
) -> tuple[list[tuple[dict[str, Any], float, float, float, float]], bool]:
    """模拟自定义布局的流式排版，返回元素落点与是否溢出正文区。首次使用：S3。

    渲染器和确定性评审共用这个函数，保证“画出来会溢出”和“检查会报溢出”用的是同一套几何。
    """

    gap = 0.15
    area_left, area_top, area_width, area_bottom = _custom_body_area()
    area_right = area_left + area_width
    x, y, row_height = area_left, area_top, 0.0
    first = True
    placements: list[tuple[dict[str, Any], float, float, float, float]] = []
    overflow = False
    for element in elements:
        cols = int(element.get("cols", 12))
        width = area_width * cols / 12.0 - gap
        height = _element_height(element, width, theme)
        if not first and x + width > area_right:
            x, y = area_left, y + row_height + gap
            row_height = 0.0
        if y + height > area_bottom:
            overflow = True
            break
        placements.append((element, x, y, width, height))
        x += width + gap
        row_height = max(row_height, height)
        first = False
    return placements, overflow


# 版式名到渲染函数的映射；名称词表由 core/contracts.py 的 KNOWN_LAYOUTS 提供。
LAYOUT_RENDERERS = {
    "bullets": _render_bullets,
    "two_column": _render_two_column,
    "comparison_table": _render_comparison_table,
    "chart_focus": _render_chart_focus,
    "custom": _render_custom,
}


def create_ppt(deck: dict[str, Any], output_path: Path, config: AppConfig) -> dict[str, Any]:
    """把已校验的 `DeckSpec` 转换为 PPTX，并按每页 layout 分派渲染。首次使用：S2。"""

    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise PresentationToolError("缺少 python-pptx；请安装 src/requirements.txt") from exc

    slide_count = int(config.section("task")["slide_count"])
    layouts = config.layout_catalog()
    # 校验时用完整词表：即使某版式未在当前模式开放，历史 Deck 也能被重新渲染。
    validate_deck(deck, slide_count)
    presentation_config = config.section("presentation")
    template = config.resolve_path(presentation_config.get("template"))
    if template and not template.exists():
        raise PresentationToolError(f"PPT 模板不存在：{template}")
    prs = Presentation(str(template)) if template else Presentation()
    if len(prs.slides):
        raise PresentationToolError("模板必须是不含现有页面的空白 .pptx")
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    theme = _theme_settings(config)
    blank_layout = prs.slide_layouts[6]
    stat: dict[str, Any] = {"charts": 0, "tables": 0, "shrunk_here": False, "custom_overflow": False}
    shrunk_pages: list[str] = []
    overflow_pages: list[str] = []

    for page_number, slide_data in enumerate(deck["slides"], start=1):
        slide = prs.slides.add_slide(blank_layout)
        stat["shrunk_here"] = False
        stat["custom_overflow"] = False
        _draw_title(slide, slide_data, theme)
        renderer = LAYOUT_RENDERERS.get(str(slide_data.get("layout", "")), _render_bullets)
        renderer(slide, slide_data, theme, stat)
        if stat["shrunk_here"]:
            shrunk_pages.append(slide_data["id"])
        if stat["custom_overflow"]:
            overflow_pages.append(slide_data["id"])
        source_text = "来源：" + ("；".join(slide_data["source_ids"]) or "本页无外部来源")
        _add_textbox(
            slide,
            source_text,
            0.7,
            6.72,
            11.6,
            0.35,
            theme["font_family"],
            int(theme["small_size"]),
            color=theme["muted_color"],
        )
        _add_textbox(
            slide,
            str(page_number),
            12.45,
            6.72,
            0.3,
            0.3,
            theme["font_family"],
            int(theme["small_size"]),
            color=theme["accent"],
            bold=True,
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return {
        "pptx_path": str(output_path),
        "slide_count": len(prs.slides),
        "chart_count": stat["charts"],
        "table_count": stat["tables"],
        "shrunk_font_pages": shrunk_pages,
        "overflow_pages": overflow_pages,
        "layouts": layouts,
    }


def render_ppt(pptx_path: Path, output_dir: Path, dpi: int) -> dict[str, Any]:
    """在工具可用时通过 LibreOffice 和 pdftoppm 把 PPT 渲染为图片。首次使用：S7。"""

    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    pdftoppm = shutil.which("pdftoppm")
    if not libreoffice or not pdftoppm:
        return {
            "rendered": False,
            "images": [],
            "warning": "未找到 LibreOffice 或 pdftoppm；已降级为结构检查",
        }
    output_dir.mkdir(parents=True, exist_ok=True)
    conversion = subprocess.run(
        [libreoffice, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(pptx_path)],
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    if conversion.returncode != 0:
        raise PresentationToolError(f"LibreOffice 渲染失败：{conversion.stderr.strip()}")
    pdf_path = output_dir / f"{pptx_path.stem}.pdf"
    if not pdf_path.exists():
        raise PresentationToolError("LibreOffice 未生成预期 PDF")
    raster = subprocess.run(
        [pdftoppm, "-png", "-r", str(dpi), str(pdf_path), str(output_dir / "slide")],
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    if raster.returncode != 0:
        raise PresentationToolError(f"pdftoppm 转图失败：{raster.stderr.strip()}")
    images = sorted(str(path) for path in output_dir.glob("slide-*.png"))
    return {"rendered": True, "pdf_path": str(pdf_path), "images": images}


def patch_deck(deck: dict[str, Any], patch_data: dict[str, Any]) -> dict[str, Any]:
    """在不直接编辑 PPTX 的情况下应用结构化整页修订。首次使用：S7。"""

    return apply_deck_patches(deck, patch_data)
