"""对 DeckSpec 内容和 PPT 元素边界执行确定性检查。首次使用：S7。"""

from __future__ import annotations

from pathlib import Path
import re
from typing import Any

from core.contracts import KNOWN_LAYOUTS
from core.config import AppConfig
from tools.presentation import _theme_settings, simulate_custom_flow


NUMERIC_PATTERN = re.compile(r"\d+(?:[.,]\d+)?(?:万|亿|%|人|次)?")

# 允许“有数据但不画图”的显式说法：必须让读者看到缺口声明，而不是直接缺图。
GAP_MARKERS = ("缺口", "缺失", "未获得", "未披露", "无法", "无同口径", "不足三", "不足 3")

# 大纲里定义要呈现数据的页面：s2 是同口径 MAU 与样本内份额，s3 是多期趋势。
# 其余页面是叙述页，提到“份额”通常只是口径说明，不应强制配图。
CHART_REQUIRED_SLIDES = {"s2", "s3"}


def _issue(
    issue_id: str,
    slide_id: str,
    issue_type: str,
    severity: str,
    evidence: str,
    suggestion: str,
) -> dict[str, str]:
    """构造一条字段稳定的 `ReviewIssue`。首次使用：S7。"""

    return {
        "id": issue_id,
        "slide_id": slide_id,
        "type": issue_type,
        "severity": severity,
        "evidence": evidence,
        "suggestion": suggestion,
    }


def check_deck(
    deck: dict[str, Any],
    known_source_ids: set[str],
    config: AppConfig,
    pptx_path: Path | None = None,
) -> list[dict[str, str]]:
    """检查文字密度、指标标签、来源引用和可选的 PPT 边界。首次使用：S7。"""

    review = config.section("review")
    max_bullets = int(review.get("max_bullets_per_slide", 6))
    max_chars = int(review.get("max_chars_per_slide", 260))
    max_table_rows = int(review.get("max_table_rows", 8))
    max_cell_chars = int(review.get("max_chars_per_table_cell", 40))
    catalog = config.layout_catalog()
    issues: list[dict[str, str]] = []
    counter = 1
    for slide in deck["slides"]:
        slide_id = slide["id"]
        if len(slide["bullets"]) > max_bullets:
            issues.append(
                _issue(f"issue_{counter}", slide_id, "too_many_bullets", "warning", str(len(slide["bullets"])), f"不超过 {max_bullets} 条")
            )
            counter += 1
        char_count = sum(len(item) for item in slide["bullets"])
        if char_count > max_chars:
            issues.append(
                _issue(f"issue_{counter}", slide_id, "too_much_text", "warning", str(char_count), f"压缩到 {max_chars} 字以内")
            )
            counter += 1
        unknown = sorted(set(slide["source_ids"]) - known_source_ids)
        if unknown:
            issues.append(
                _issue(f"issue_{counter}", slide_id, "unknown_source", "error", ", ".join(unknown), "删除或替换未知来源 ID")
            )
            counter += 1
        joined = " ".join(slide["bullets"])
        layout = str(slide.get("layout", ""))
        if layout not in KNOWN_LAYOUTS or layout not in catalog:
            issues.append(
                _issue(
                    f"issue_{counter}",
                    slide_id,
                    "unknown_layout",
                    "error",
                    layout or "缺失",
                    f"改用当前模式允许的版式：{catalog}",
                )
            )
            counter += 1
        table = slide.get("table") or {}
        table_columns = table.get("columns") or []
        table_rows = table.get("rows") or []
        if layout == "comparison_table" and not table_rows:
            issues.append(
                _issue(
                    f"issue_{counter}",
                    slide_id,
                    "table_missing",
                    "error",
                    "comparison_table 版式没有表格数据",
                    "补上 table 的 columns 与 rows，或改用 bullets 版式",
                )
            )
            counter += 1
        if layout == "custom" and slide.get("elements"):
            _, overflow = simulate_custom_flow(slide["elements"], _theme_settings(config))
            if overflow:
                issues.append(
                    _issue(
                        f"issue_{counter}",
                        slide_id,
                        "custom_layout_overflow",
                        "error",
                        "自定义布局内容超出正文区，与来源行重叠",
                        "减少元素、缩小 cols 或删减文字",
                    )
                )
                counter += 1
        if len(table_rows) > max_table_rows:
            issues.append(
                _issue(
                    f"issue_{counter}",
                    slide_id,
                    "too_many_table_rows",
                    "warning",
                    str(len(table_rows)),
                    f"表格数据行不超过 {max_table_rows} 行",
                )
            )
            counter += 1
        long_cells = [cell for row in table_rows for cell in row if len(str(cell)) > max_cell_chars]
        if long_cells:
            issues.append(
                _issue(
                    f"issue_{counter}",
                    slide_id,
                    "table_cell_too_long",
                    "warning",
                    str(long_cells[0])[:60],
                    f"单元格文字压缩到 {max_cell_chars} 字以内",
                )
            )
            counter += 1
        if review.get("require_source_for_numeric_claim", True) and NUMERIC_PATTERN.search(joined) and not slide["source_ids"]:
            issues.append(
                _issue(f"issue_{counter}", slide_id, "numeric_claim_without_source", "error", joined[:160], "为数字添加已读取来源，或删除无依据数字")
            )
            counter += 1
        if (
            review.get("require_metric_scope_label", True)
            and slide_id == "s2"
            and "份额" in joined
            and "样本内 MAU 份额" not in joined
        ):
            issues.append(
                _issue(f"issue_{counter}", slide_id, "share_label", "error", "份额标签不完整", "统一写为“样本内 MAU 份额”")
            )
            counter += 1
        chart = slide.get("chart") or {}
        chart_type = str(chart.get("type", "none"))
        if chart_type in {"bar", "line"}:
            # 有图时先确认数据自洽，避免画出空图或长度错位的图。
            categories = chart.get("categories") or []
            series = chart.get("series") or []
            if not categories or not series:
                issues.append(
                    _issue(f"issue_{counter}", slide_id, "invalid_chart", "error", chart_type, "补齐 categories 与 series，或改为 type=none")
                )
                counter += 1
            elif any(len(item.get("values", [])) != len(categories) for item in series):
                issues.append(
                    _issue(f"issue_{counter}", slide_id, "invalid_chart", "error", "数值与分类数量不一致", "让每个 series 的 values 与 categories 等长")
                )
                counter += 1
            if not slide["source_ids"]:
                issues.append(
                    _issue(f"issue_{counter}", slide_id, "chart_without_source", "error", "图表无来源", "为图表补上已读取的来源，或删除图表")
                )
                counter += 1
        elif review.get("require_chart_for_share", True) and slide_id in CHART_REQUIRED_SLIDES:
            # 这两页天然需要图；实在没有同口径数据时必须写明缺口，而不是直接缺图。
            if not any(marker in joined for marker in GAP_MARKERS):
                issues.append(
                    _issue(
                        f"issue_{counter}",
                        slide_id,
                        "chart_missing",
                        "warning",
                        "该页涉及份额或多期数值但没有图表",
                        "补充同口径图表，或明确写出数据缺口",
                    )
                )
                counter += 1
    if pptx_path and pptx_path.exists():
        for bounds_issue in check_ppt_bounds(pptx_path, counter):
            issues.append(bounds_issue)
            counter += 1
    return issues


def review_deck_with_model(model: Any, deck: dict[str, Any], config: AppConfig) -> list[dict[str, str]]:
    """让模型对 DeckSpec 做一轮语义评审，返回带 source=model 的 issue 列表。首次使用：S7。"""

    from core.contracts import review_issue_schema
    from core.prompts import deck_review_prompt

    max_issues = int(config.section("review").get("max_model_issues", 5))
    generated = model.generate_structured(
        deck_review_prompt(config, deck),
        "只提确有把握的问题，输出 JSON。",
        "deck_review",
        review_issue_schema(max_issues),
    )
    valid_slides = {f"s{i}" for i in range(1, int(config.section("task")["slide_count"]) + 1)}
    issues: list[dict[str, str]] = []
    for item in generated.get("issues", []):
        if item.get("slide_id") not in valid_slides:
            continue
        issue = dict(item)
        issue["source"] = "model"
        issues.append(issue)
    return issues[:max_issues]


def check_ppt_bounds(pptx_path: Path, starting_counter: int = 1) -> list[dict[str, str]]:
    """标记超出幻灯片物理画布边界的元素。首次使用：S7。"""

    try:
        from pptx import Presentation
    except ImportError:
        return [
            _issue(
                f"issue_{starting_counter}",
                "deck",
                "dependency_missing",
                "warning",
                "python-pptx unavailable",
                "安装 src/requirements.txt 后重新检查",
            )
        ]
    prs = Presentation(str(pptx_path))
    issues: list[dict[str, str]] = []
    counter = starting_counter
    for page_number, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width or shape.top + shape.height > prs.slide_height:
                issues.append(
                    _issue(
                        f"issue_{counter}",
                        f"s{page_number}",
                        "shape_out_of_bounds",
                        "error",
                        getattr(shape, "name", "unknown shape"),
                        "缩小或移动该元素",
                    )
                )
                counter += 1
    return issues
